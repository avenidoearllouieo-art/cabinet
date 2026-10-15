import re
import zipfile

from django.core.exceptions import ValidationError
from django.core.validators import validate_email
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from .models import Section, User


IMPORT_COLUMNS = ('student_id', 'first_name', 'last_name', 'email', 'role', 'subject')
EXPECTED_HEADERS = ('student id', 'first name', 'last name', 'email', 'role', 'subject')
MAX_IMPORT_ROWS = 1000


def _cell_text(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _row_error(row_number, column, message):
    return {'row': row_number, 'column': column, 'message': message}


def read_student_workbook(uploaded_file):
    try:
        workbook = load_workbook(uploaded_file, read_only=True, data_only=True)
    except (InvalidFileException, OSError, ValueError, zipfile.BadZipFile):
        return [], [_row_error(1, 'file', 'The uploaded file is not a readable .xlsx workbook.')]

    try:
        worksheet = workbook.active
        iterator = worksheet.iter_rows(values_only=True)
        header_row = next(iterator, None)
        if not header_row:
            return [], [_row_error(1, 'header', 'The workbook is empty.')]

        headers = [_cell_text(value).casefold().replace('_', ' ') for value in header_row]
        while headers and not headers[-1]:
            headers.pop()
        normalized_headers = tuple(' '.join(header.split()) for header in headers)
        if normalized_headers != EXPECTED_HEADERS:
            expected = ' | '.join(('Student ID', 'First Name', 'Last Name', 'Email', 'Role', 'Subject'))
            return [], [_row_error(
                1,
                'header',
                f'Expected columns in this exact order: {expected}. Download the current template; older templates are not supported.',
            )]

        errors = []
        rows = []
        for row_number, values in enumerate(iterator, start=2):
            if not any(_cell_text(value) for value in values):
                continue
            if any(_cell_text(value) for value in values[len(IMPORT_COLUMNS):]):
                errors.append(_row_error(row_number, 'row', 'Unexpected values found after the Subject column.'))
                continue
            if len(rows) >= MAX_IMPORT_ROWS:
                errors.append(_row_error(row_number, 'row', f'Files may contain at most {MAX_IMPORT_ROWS} student rows.'))
                break
            row = {
                column: _cell_text(values[index] if index < len(values) else None)
                for index, column in enumerate(IMPORT_COLUMNS)
            }
            row['row_number'] = row_number
            rows.append(row)

        if not rows and not errors:
            errors.append(_row_error(2, 'row', 'The workbook contains no student rows.'))
        return rows, errors
    finally:
        workbook.close()


def username_for_student_id(student_id):
    safe_id = re.sub(r'[^A-Za-z0-9@.+_-]+', '_', student_id).strip('_')
    return f'student_{safe_id or "record"}'


def build_student(row):
    section = Section.objects.get(pk=row['section_id']) if row['section_id'] else None
    user = User(
        username=row['username'],
        first_name=row['first_name'],
        last_name=row['last_name'],
        student_id=row['student_id'],
        email=row['email'],
        role=User.RoleChoices.STUDENT,
        section=section,
        nfc_uid=None,
        is_active=row['active'],
        is_staff=False,
        is_superuser=False,
    )
    user.set_unusable_password()
    return user


def validate_student_rows(rows):
    errors = []
    prepared = []
    seen = {key: {} for key in ('student_id', 'email', 'username')}
    section_lookup = {}
    for section in Section.objects.all().only('pk', 'subject_code'):
        section_lookup.setdefault(section.subject_code.strip().casefold(), []).append(section.pk)

    for source in rows:
        row_number = int(source.get('row_number') or 0)
        row_errors = []
        student_id = _cell_text(source.get('student_id'))
        first_name = _cell_text(source.get('first_name'))
        last_name = _cell_text(source.get('last_name'))
        email = _cell_text(source.get('email'))
        subject = _cell_text(source.get('subject'))
        username = _cell_text(source.get('username')) or username_for_student_id(student_id)
        role = _cell_text(source.get('role')).casefold()

        if role and role != User.RoleChoices.STUDENT:
            row_errors.append(_row_error(row_number, 'role', 'Only the student role can be imported here.'))

        for column, value in (
            ('student_id', student_id),
            ('first_name', first_name),
            ('last_name', last_name),
            ('email', email),
            ('role', role),
        ):
            if not value:
                row_errors.append(_row_error(row_number, column, f'{column} is required.'))

        if student_id and len(student_id) > User._meta.get_field('student_id').max_length:
            row_errors.append(_row_error(row_number, 'student_id', 'Student ID is too long.'))
        if email:
            try:
                validate_email(email)
            except ValidationError:
                row_errors.append(_row_error(row_number, 'email', 'Enter a valid email address.'))

        for column, value, normalized in (
            ('student_id', student_id, student_id.casefold()),
            ('email', email, email.casefold()),
            ('username', username, username.casefold()),
        ):
            if not value:
                continue
            if normalized in seen[column]:
                row_errors.append(_row_error(row_number, column, f'Duplicate {column} in this spreadsheet (first seen on row {seen[column][normalized]}).'))
            else:
                seen[column][normalized] = row_number

        if student_id and User.objects.filter(student_id=student_id).exists():
            row_errors.append(_row_error(row_number, 'student_id', 'A student with this ID already exists; imports do not overwrite records.'))
        if email and User.objects.filter(email__iexact=email).exists():
            row_errors.append(_row_error(row_number, 'email', 'This email address is already in use.'))
        if username and User.objects.filter(username=username).exists():
            row_errors.append(_row_error(row_number, 'username', 'This username is already in use.'))

        section_id = None
        if subject:
            matches = section_lookup.get(subject.casefold(), [])
            if not matches:
                row_errors.append(_row_error(row_number, 'subject', 'No existing Section has this subject code.'))
            elif len(matches) > 1:
                row_errors.append(_row_error(row_number, 'subject', 'More than one Section has this subject code; resolve the duplicate Sections first.'))
            else:
                section_id = matches[0]

        candidate = {
            'row_number': row_number,
            'student_id': student_id,
            'first_name': first_name,
            'last_name': last_name,
            'email': email,
            'username': username,
            'subject': subject,
            'role': role,
            'section_id': section_id,
            'active': True,
        }

        if not row_errors:
            try:
                build_student(candidate).full_clean()
            except ValidationError as error:
                if hasattr(error, 'message_dict'):
                    for field, messages in error.message_dict.items():
                        for message in messages:
                            row_errors.append(_row_error(row_number, field, str(message)))
                else:
                    row_errors.append(_row_error(row_number, 'row', '; '.join(error.messages)))

        errors.extend(row_errors)
        prepared.append(candidate)

    return prepared, errors