from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.exceptions import ValidationError
from django.contrib import admin as django_admin
from django.test import Client, RequestFactory, TestCase, override_settings
from django.test import TransactionTestCase
from django.db import connection
from django.db.migrations.executor import MigrationExecutor
from django.urls import reverse
from django.utils import timezone
from rest_framework.test import APIClient

from io import BytesIO
from PIL import Image
import os
import tempfile
from datetime import timedelta
from urllib.parse import parse_qs, urlparse

from openpyxl import Workbook, load_workbook

from .admin import AccessLogAdmin
from .models import AccessLog, AccessLogPhoto, CabinetEvent, Notification, Section, User, Activity, ActivityAttachment, Submission, PasswordResetRequest, NFCEnrollmentSession, CabinetSession


class InstructorSectionOverviewScopeTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.instructor = User.objects.create_user(
            username='section-overview-instructor',
            email='section-overview-instructor@example.com',
            password='secret1234',
            role=User.RoleChoices.INSTRUCTOR,
            instructor_id='OVERVIEW-INST-1',
        )
        self.section_a = Section.objects.create(
            section_name='Overview Section A',
            subject_code='OV-A',
            program='Program A',
            year_level='3',
            academic_year='2026-2027',
            instructor=self.instructor,
        )
        self.section_b = Section.objects.create(
            section_name='Overview Section B',
            subject_code='OV-B',
            program='Program B',
            year_level='2',
            academic_year='2026-2027',
        )
        self.section_b.assigned_instructors.add(self.instructor)
        self.student_a = User.objects.create_user(
            username='overview-student-a',
            email='overview-student-a@example.com',
            password='secret1234',
            role=User.RoleChoices.STUDENT,
            student_id='OV-STU-A',
            section=self.section_a,
        )
        self.student_b = User.objects.create_user(
            username='overview-student-b',
            email='overview-student-b@example.com',
            password='secret1234',
            role=User.RoleChoices.STUDENT,
            student_id='OV-STU-B',
            section=self.section_b,
        )
        self.student_b_missing = User.objects.create_user(
            username='overview-student-b-missing',
            email='overview-student-b-missing@example.com',
            password='secret1234',
            role=User.RoleChoices.STUDENT,
            student_id='OV-STU-B-MISSING',
            section=self.section_b,
        )
        self.activity_a = Activity.objects.create(
            title='Section A activity',
            created_by=self.instructor,
            cabinet_station='Cabinet Station 1',
        )
        self.activity_a.assigned_sections.add(self.section_a)
        self.activity_b = Activity.objects.create(
            title='Section B activity',
            created_by=self.instructor,
            cabinet_station='Cabinet Station 2',
        )
        self.activity_b.assigned_sections.add(self.section_b)
        Submission.objects.create(activity=self.activity_a, student=self.student_a)
        Submission.objects.create(activity=self.activity_b, student=self.student_b)
        AccessLog.objects.create(user=self.student_a, station='Cabinet Station 1', status='success')
        AccessLog.objects.create(user=self.student_b, station='Cabinet Station 2', status='failed')
        self.client.force_authenticate(user=self.instructor)

    def test_section_overview_contains_only_selected_sections_records(self):
        response = self.client.get(f'/api/sections/{self.section_a.pk}/overview/')

        self.assertEqual(response.status_code, 200, response.json())
        payload = response.json()
        self.assertEqual(payload['section']['section_id'], self.section_a.pk)
        self.assertEqual(payload['counts'], {
            'students': 1,
            'activities': 1,
            'submitted': 1,
            'expected': 1,
            'missing': 0,
            'late': 0,
            'pending_review': 1,
        })
        self.assertEqual({student['id'] for student in payload['students']}, {self.student_a.pk})
        self.assertEqual({activity['id'] for activity in payload['activities']}, {self.activity_a.pk})
        self.assertEqual({log['user'] for log in payload['access_logs']}, {self.student_a.pk})
        self.assertEqual(payload['cabinet_stations'], [{'station': 'Cabinet Station 1', 'status': 'Available'}])

    def test_assigned_sections_and_section_list_metrics_respect_selected_section(self):
        response = self.client.get('/api/sections/')

        self.assertEqual(response.status_code, 200, response.json())
        sections = response.json()['results']
        self.assertEqual({section['section_id'] for section in sections}, {self.section_a.pk, self.section_b.pk})
        by_id = {section['section_id']: section for section in sections}
        self.assertEqual(by_id[self.section_a.pk]['student_count'], 1)
        self.assertEqual(by_id[self.section_a.pk]['activity_count'], 1)
        self.assertEqual(by_id[self.section_a.pk]['submitted_count'], 1)
        self.assertEqual(by_id[self.section_a.pk]['missing_count'], 0)
        self.assertEqual(by_id[self.section_a.pk]['cabinet_station'], 'Cabinet Station 1')
        self.assertEqual(by_id[self.section_b.pk]['student_count'], 2)
        self.assertEqual(
            by_id[self.section_b.pk]['instructor_name'],
            self.instructor.get_full_name().strip() or self.instructor.username,
        )
        self.assertEqual(by_id[self.section_b.pk]['activity_count'], 1)
        self.assertEqual(by_id[self.section_b.pk]['submitted_count'], 1)
        self.assertEqual(by_id[self.section_b.pk]['expected_submission_count'], 2)
        self.assertEqual(by_id[self.section_b.pk]['missing_count'], 1)
        self.assertEqual(by_id[self.section_b.pk]['cabinet_station'], 'Cabinet Station 2')

        students_response = self.client.get(f'/api/users/?section={self.section_b.pk}&role=student')
        self.assertEqual(students_response.status_code, 200, students_response.json())
        self.assertEqual(
            {student['id'] for student in students_response.json()['results']},
            {self.student_b.pk, self.student_b_missing.pk},
        )

    def test_section_destination_endpoints_are_scoped_to_selected_section(self):
        activities_response = self.client.get(f'/api/activities/?section={self.section_a.pk}')
        submissions_response = self.client.get(f'/api/submissions/?section={self.section_a.pk}')
        logs_response = self.client.get(f'/api/access-logs/?user__section={self.section_a.pk}')
        stats_response = self.client.get(f'/api/access-logs/stats/?user__section={self.section_a.pk}')

        self.assertEqual(activities_response.status_code, 200, activities_response.json())
        self.assertEqual(submissions_response.status_code, 200, submissions_response.json())
        self.assertEqual(logs_response.status_code, 200, logs_response.json())
        self.assertEqual(stats_response.status_code, 200, stats_response.json())
        self.assertEqual({activity['id'] for activity in activities_response.json()['results']}, {self.activity_a.pk})
        self.assertEqual({submission['id'] for submission in submissions_response.json()['results']}, {self.student_a.submissions.get().pk})
        self.assertEqual({log['user'] for log in logs_response.json()['results']}, {self.student_a.pk})
        self.assertEqual(stats_response.json()['total_accesses_today'], 1)

    def test_instructor_cannot_open_an_unassigned_section_overview(self):
        other = User.objects.create_user(
            username='other-section-instructor',
            email='other-section-instructor@example.com',
            password='secret1234',
            role=User.RoleChoices.INSTRUCTOR,
            instructor_id='OVERVIEW-INST-2',
        )
        other_section = Section.objects.create(section_name='Other Section', instructor=other)

        response = self.client.get(f'/api/sections/{other_section.pk}/overview/')

        self.assertEqual(response.status_code, 404)


class InstructorSectionAssignmentTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.admin = User.objects.create_user(
            username='api-admin',
            email='api-admin@example.com',
            password='secret1234',
            role=User.RoleChoices.ADMIN,
        )
        self.first_instructor = User.objects.create_user(
            username='first-instructor',
            email='first-instructor@example.com',
            password='secret1234',
            role=User.RoleChoices.INSTRUCTOR,
            instructor_id='INST-001',
        )
        self.second_instructor = User.objects.create_user(
            username='second-instructor',
            email='second-instructor@example.com',
            password='secret1234',
            role=User.RoleChoices.INSTRUCTOR,
            instructor_id='INST-002',
        )
        self.section_a = Section.objects.create(section_name='API Section A', subject_code='API-A')
        self.section_b = Section.objects.create(section_name='API Section B', subject_code='API-B')
        self.client.force_authenticate(user=self.admin)

    def test_create_instructor_accepts_ids_and_returns_section_objects(self):
        response = self.client.post('/api/users/', {
            'username': 'created-instructor',
            'email': 'created-instructor@example.com',
            'password': 'secret1234',
            'role': User.RoleChoices.INSTRUCTOR,
            'instructor_id': 'INST-003',
            'assigned_sections': [self.section_a.pk, self.section_b.pk],
        }, format='json')

        self.assertEqual(response.status_code, 201, response.json())
        created = User.objects.get(pk=response.json()['id'])
        self.assertEqual(
            {item['section_id'] for item in response.json()['assigned_sections']},
            {self.section_a.pk, self.section_b.pk},
        )
        self.assertEqual(set(created.assigned_sections.values_list('pk', flat=True)), {self.section_a.pk, self.section_b.pk})
        self.assertFalse(created.check_password('created-instructor'))
        self.assertTrue(created.check_password('secret1234'))
        self.assertFalse(Section.objects.filter(pk__in=[self.section_a.pk, self.section_b.pk]).exclude(instructor=created).exists())

    def test_assignment_transfer_removes_previous_owner(self):
        self.first_instructor.assigned_sections.add(self.section_a)
        self.section_a.instructor = self.first_instructor
        self.section_a.save(update_fields=['instructor'])

        response = self.client.patch(f'/api/users/{self.second_instructor.pk}/', {
            'assigned_sections': [self.section_a.pk],
        }, format='json')

        self.assertEqual(response.status_code, 200, response.json())
        self.section_a.refresh_from_db()
        self.assertEqual(self.section_a.instructor, self.second_instructor)
        self.assertFalse(self.first_instructor.assigned_sections.filter(pk=self.section_a.pk).exists())
        self.assertTrue(self.second_instructor.assigned_sections.filter(pk=self.section_a.pk).exists())

    def test_replacing_and_clearing_assignments_updates_both_relationships(self):
        self.first_instructor.assigned_sections.add(self.section_a)
        self.section_a.instructor = self.first_instructor
        self.section_a.save(update_fields=['instructor'])

        response = self.client.patch(f'/api/users/{self.first_instructor.pk}/', {
            'assigned_sections': [self.section_b.pk],
        }, format='json')
        self.assertEqual(response.status_code, 200, response.json())
        self.section_a.refresh_from_db()
        self.section_b.refresh_from_db()
        self.assertIsNone(self.section_a.instructor)
        self.assertEqual(self.section_b.instructor, self.first_instructor)

        response = self.client.patch(f'/api/users/{self.first_instructor.pk}/', {
            'assigned_sections': [],
        }, format='json')
        self.assertEqual(response.status_code, 200, response.json())
        self.section_b.refresh_from_db()
        self.assertIsNone(self.section_b.instructor)
        self.assertFalse(self.first_instructor.assigned_sections.exists())

    def test_role_change_clears_instructor_assignments(self):
        self.first_instructor.assigned_sections.add(self.section_a)
        self.section_a.instructor = self.first_instructor
        self.section_a.save(update_fields=['instructor'])

        response = self.client.patch(f'/api/users/{self.first_instructor.pk}/', {
            'role': User.RoleChoices.ADMIN,
        }, format='json')

        self.assertEqual(response.status_code, 200, response.json())
        self.section_a.refresh_from_db()
        self.first_instructor.refresh_from_db()
        self.assertIsNone(self.section_a.instructor)
        self.assertFalse(self.first_instructor.assigned_sections.exists())
        self.assertEqual(response.json()['assigned_sections'], [])

    def test_section_endpoint_keeps_assignment_relationships_in_sync(self):
        response = self.client.patch(f'/api/sections/{self.section_a.pk}/', {
            'instructor': self.first_instructor.pk,
        }, format='json')
        self.assertEqual(response.status_code, 200, response.json())
        self.assertTrue(self.first_instructor.assigned_sections.filter(pk=self.section_a.pk).exists())

        response = self.client.patch(f'/api/sections/{self.section_a.pk}/', {
            'instructor': self.second_instructor.pk,
        }, format='json')
        self.assertEqual(response.status_code, 200, response.json())
        self.assertFalse(self.first_instructor.assigned_sections.filter(pk=self.section_a.pk).exists())
        self.assertTrue(self.second_instructor.assigned_sections.filter(pk=self.section_a.pk).exists())

    def test_legacy_section_owner_is_included_in_user_response(self):
        self.section_a.instructor = self.first_instructor
        self.section_a.save(update_fields=['instructor'])

        response = self.client.get(f'/api/users/{self.first_instructor.pk}/')

        self.assertEqual(response.status_code, 200, response.json())
        self.assertEqual(
            [item['section_id'] for item in response.json()['assigned_sections']],
            [self.section_a.pk],
        )


class ActivityCreationTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.instructor = User.objects.create_user(
            username='instructor1',
            email='instructor1@example.com',
            password='secret123',
            role='instructor',
            first_name='Instructor',
            last_name='One',
            instructor_id='I001',
            nfc_uid='NFC-INST-001',
        )
        self.section_a = Section.objects.create(section_name='Section A', subject_code='SEC-A', instructor=self.instructor)
        self.section_b = Section.objects.create(section_name='Section B', subject_code='SEC-B', instructor=self.instructor)
        self.other_instructor = User.objects.create_user(
            username='instructor2',
            email='instructor2@example.com',
            password='secret123',
            role='instructor',
            instructor_id='I002',
            nfc_uid='NFC-INST-002',
        )

    def test_instructor_can_create_activity_with_extended_fields(self):
        self.client.force_authenticate(user=self.instructor)

        payload = {
            'title': 'Extended Activity',
            'description': 'A richer activity form payload.',
            'instructions': '<p>Read the brief carefully.</p>',
            'activity_type': 'Assignment',
            'assigned_sections': [self.section_a.section_id, self.section_b.section_id],
            'assigned_instructor': self.instructor.id,
            'due_date': '2026-08-10T09:00:00Z',
            'max_score': 100,
            'cabinet_station': 'Cabinet 1',
            'status': 'Draft',
        }

        response = self.client.post('/api/activities/', payload, format='json')

        self.assertEqual(response.status_code, 201)
        data = response.json()
        self.assertEqual(data['title'], 'Extended Activity')
        self.assertEqual(data['activity_type'], 'Assignment')
        self.assertEqual(data['status'], 'Draft')
        self.assertEqual(data['cabinet_station'], 'Cabinet 1')
        self.assertEqual(data['assigned_instructor'], self.instructor.id)
        self.assertEqual(set(data['assigned_sections']), {self.section_a.section_id, self.section_b.section_id})

        activity = Activity.objects.get(id=data['id'])
        self.assertEqual(activity.created_by, self.instructor)
        self.assertEqual(activity.instructions, '<p>Read the brief carefully.</p>')
        self.assertTrue(activity.assigned_sections.filter(section_id__in=[self.section_a.section_id, self.section_b.section_id]).count() == 2)

    def test_create_assigns_authenticated_instructor_and_list_is_isolated(self):
        self.client.force_authenticate(user=self.instructor)
        response = self.client.post('/api/activities/', {
            'title': 'Owned Activity',
            'section': self.section_a.section_id,
        }, format='json')

        self.assertEqual(response.status_code, 201)
        activity = Activity.objects.get(id=response.json()['id'])
        self.assertEqual(activity.created_by_id, self.instructor.id)

        self.client.force_authenticate(user=self.other_instructor)
        other_response = self.client.get('/api/activities/')
        other_results = other_response.json().get('results', other_response.json())
        self.assertNotIn(activity.id, [item['id'] for item in other_results])

        self.client.force_authenticate(user=self.instructor)
        own_response = self.client.get('/api/activities/')
        own_results = own_response.json().get('results', own_response.json())
        self.assertIn(activity.id, [item['id'] for item in own_results])

    def test_update_preserves_creator_and_rejects_client_ownership_change(self):
        activity = Activity.objects.create(title='Owned Activity', created_by=self.instructor)
        self.client.force_authenticate(user=self.instructor)

        response = self.client.patch(f'/api/activities/{activity.id}/', {
            'title': 'Updated Activity',
            'created_by': self.other_instructor.id,
        }, format='json')

        self.assertEqual(response.status_code, 200)
        activity.refresh_from_db()
        self.assertEqual(activity.title, 'Updated Activity')
        self.assertEqual(activity.created_by_id, self.instructor.id)

    def test_instructor_can_upload_files_with_activity(self):
        self.client.force_authenticate(user=self.instructor)
        file_upload = SimpleUploadedFile('brief.pdf', b'file-content', content_type='application/pdf')

        response = self.client.post('/api/activities/', {
            'title': 'Activity with File',
            'description': 'An activity with uploaded files.',
            'instructions': '<p>Read the brief.</p>',
            'activity_type': 'Project',
            'assigned_sections': [self.section_a.section_id],
            'assigned_instructor': self.instructor.id,
            'due_date': '2026-08-10T09:00:00Z',
            'max_score': 100,
            'cabinet_station': 'Cabinet 2',
            'status': 'Published',
            'attachments': [file_upload],
        }, format='multipart')

        self.assertEqual(response.status_code, 201)
        data = response.json()
        self.assertTrue(data['attachments'])
        self.assertEqual(data['attachments'][0]['filename'], 'brief.pdf')
        activity = Activity.objects.get(id=data['id'])
        self.assertTrue(activity.attachments.exists())

    def test_instructor_can_edit_activity_and_keep_existing_attachments(self):
        self.client.force_authenticate(user=self.instructor)
        activity = Activity.objects.create(
            title='Editable Activity',
            description='Original description',
            instructions='Original instructions',
            created_by=self.instructor,
            section=self.section_a,
        )
        existing_attachment = ActivityAttachment.objects.create(activity=activity, file=SimpleUploadedFile('old.pdf', b'old', content_type='application/pdf'))

        new_file = SimpleUploadedFile('new.pdf', b'new', content_type='application/pdf')
        response = self.client.patch(f'/api/activities/{activity.id}/', {
            'title': 'Updated Activity',
            'attachments': [new_file],
        }, format='multipart')

        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data['title'], 'Updated Activity')
        self.assertEqual(len(data['attachments']), 2)
        remaining_ids = {item['id'] for item in data['attachments']}
        self.assertIn(existing_attachment.id, remaining_ids)

    def test_delete_attachment_endpoint_removes_database_and_file(self):
        self.client.force_authenticate(user=self.instructor)
        activity = Activity.objects.create(
            title='Deleteable Activity',
            description='desc',
            instructions='instr',
            created_by=self.instructor,
            section=self.section_a,
        )
        attachment = ActivityAttachment.objects.create(activity=activity, file=SimpleUploadedFile('delete-me.pdf', b'data', content_type='application/pdf'))

        response = self.client.delete(f'/api/activities/{activity.id}/attachments/{attachment.id}/')

        self.assertEqual(response.status_code, 204)
        self.assertFalse(ActivityAttachment.objects.filter(id=attachment.id).exists())


class SubmissionGradingTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.instructor = User.objects.create_user(
            username='instructor-grade',
            email='instructor-grade@example.com',
            password='secret123',
            role='instructor',
            first_name='Instructor',
            last_name='Grade',
            instructor_id='I-GRADE',
            nfc_uid='NFC-GRADE-001',
        )
        self.section = Section.objects.create(section_name='Section Grade', subject_code='SEC-G', instructor=self.instructor)
        self.activity = Activity.objects.create(
            title='Graded Activity',
            description='An activity for grading tests',
            instructions='Do the work',
            created_by=self.instructor,
            section=self.section,
            max_score=100,
        )
        self.student = User.objects.create_user(
            username='student-grade',
            email='student-grade@example.com',
            password='secret123',
            role='student',
            first_name='Student',
            last_name='Grade',
            student_id='S-GRADE',
            section=self.section,
            nfc_uid='NFC-GRADE-002',
        )
        self.submission = Submission.objects.create(
            activity=self.activity,
            student=self.student,
            remarks='Initial submission',
            score=None,
            feedback='',
        )

    def test_instructor_can_grade_existing_submission(self):
        self.client.force_authenticate(user=self.instructor)

        response = self.client.post(f'/api/submissions/{self.submission.id}/grade/', {
            'score': 88,
            'feedback': 'Great work',
        }, format='json')

        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data['score'], 88)
        self.assertEqual(data['feedback'], 'Great work')

        self.submission.refresh_from_db()
        self.assertEqual(self.submission.score, 88)
        self.assertEqual(self.submission.feedback, 'Great work')
        self.assertEqual(self.submission.graded_by, self.instructor)
        self.assertIsNotNone(self.submission.graded_at)
        self.assertEqual(self.submission.status, 'Graded')

    def test_grade_rejects_score_above_maximum(self):
        self.client.force_authenticate(user=self.instructor)

        response = self.client.post(f'/api/submissions/{self.submission.id}/grade/', {
            'score': 101,
            'feedback': 'Too high',
        }, format='json')

        self.assertEqual(response.status_code, 400)
        self.assertIn('cannot exceed', response.json()['detail'])

        self.submission.refresh_from_db()
        self.assertIsNone(self.submission.score)
        self.assertEqual(self.submission.feedback, '')

    def test_grade_rejects_negative_score(self):
        self.client.force_authenticate(user=self.instructor)

        response = self.client.post(f'/api/submissions/{self.submission.id}/grade/', {
            'score': -1,
            'feedback': 'Negative score',
        }, format='json')

        self.assertEqual(response.status_code, 400)
        self.assertIn('cannot be negative', response.json()['detail'])

        self.submission.refresh_from_db()
        self.assertIsNone(self.submission.score)
        self.assertEqual(self.submission.feedback, '')


class StudentNotificationTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.student = User.objects.create_user(
            username='student3',
            email='student3@example.com',
            password='secret123',
            role='student',
            first_name='Student',
            last_name='Three',
            student_id='S003',
        )
        self.other_student = User.objects.create_user(
            username='student4',
            email='student4@example.com',
            password='secret123',
            role='student',
            first_name='Student',
            last_name='Four',
            student_id='S004',
        )
        Notification.objects.create(
            student=self.student,
            title='New activity assigned',
            message='A new activity is ready for you.',
            notification_type='deadline',
            link='/student/activities',
        )
        Notification.objects.create(
            student=self.other_student,
            title='Hidden notification',
            message='This should not be visible.',
            notification_type='submission',
            link='/student/submissions',
        )

    def test_student_only_sees_their_own_notifications(self):
        self.client.force_authenticate(user=self.student)

        response = self.client.get('/api/notifications/')

        self.assertEqual(response.status_code, 200)
        data = response.json()
        results = data.get('results', [])
        self.assertEqual(len(results), 1)
        self.assertEqual(results[0]['title'], 'New activity assigned')

    def test_student_unread_count_is_private(self):
        self.client.force_authenticate(user=self.student)

        response = self.client.get('/api/notifications/unread_count/')

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()['unread_count'], 1)

    def test_authenticated_instructor_comment_notifies_assigned_student(self):
        instructor = User.objects.create_user(
            username='comment-instructor',
            email='comment-instructor@example.com',
            password='secret123',
            role='instructor',
            instructor_id='I-COMMENT',
            nfc_uid='NFC-COMMENT',
        )
        section = Section.objects.create(section_name='Comment Section', instructor=instructor)
        self.student.section = section
        self.student.save(update_fields=['section'])
        activity = Activity.objects.create(title='Discuss this', created_by=instructor, section=section)

        self.client.force_authenticate(user=instructor)
        response = self.client.post('/api/activity-discussions/', {
            'activity': activity.id,
            'message': 'Please review this activity.',
        }, format='json')

        self.assertEqual(response.status_code, 201)
        self.client.force_authenticate(user=self.student)
        notifications = self.client.get('/api/notifications/').json()['results']
        feedback_notifications = [
            item for item in notifications
            if item['notification_type'] == 'feedback'
        ]
        self.assertEqual(len(feedback_notifications), 1)
        self.assertEqual(feedback_notifications[0]['student'], self.student.id)
        self.assertFalse(
            Notification.objects.filter(
                student=self.other_student,
                notification_type=Notification.TypeChoices.FEEDBACK,
            ).exists()
        )


class StudentProfileTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.student = User.objects.create_user(
            username='student5',
            email='student5@example.com',
            password='secret123',
            role='student',
            first_name='Student',
            last_name='Five',
            student_id='S005',
        )

    def test_student_can_view_own_profile(self):
        self.client.force_authenticate(user=self.student)

        response = self.client.get('/api/users/profile/')

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()['id'], self.student.id)
        self.assertEqual(response.json()['student_id'], 'S005')

    def test_student_profile_resolves_instructors_from_section_relationships(self):
        primary = User.objects.create_user(
            username='section-primary-instructor',
            email='section-primary-instructor@example.com',
            password='secret1234',
            role=User.RoleChoices.INSTRUCTOR,
            instructor_id='SECTION-PRIMARY',
            first_name='Primary',
            last_name='Teacher',
            contact_number='555-0101',
        )
        assigned = User.objects.create_user(
            username='section-assigned-instructor',
            email='section-assigned-instructor@example.com',
            password='secret1234',
            role=User.RoleChoices.INSTRUCTOR,
            instructor_id='SECTION-ASSIGNED',
            first_name='Assigned',
            last_name='Teacher',
            contact_number='555-0102',
        )
        section = Section.objects.create(
            section_name='Student Profile Section',
            program='Information Technology',
            year_level='3',
            academic_year='2026-2027',
            instructor=primary,
        )
        section.assigned_instructors.add(primary, assigned)
        self.student.section = section
        self.student.save(update_fields=['section'])
        self.client.force_authenticate(user=self.student)

        response = self.client.get('/api/users/profile/')

        self.assertEqual(response.status_code, 200, response.json())
        payload = response.json()
        self.assertEqual(payload['section_program'], 'Information Technology')
        self.assertEqual(payload['section_name'], 'Student Profile Section')
        self.assertEqual(
            [(instructor['id'], instructor['full_name']) for instructor in payload['section_instructors']],
            [(primary.pk, 'Primary Teacher'), (assigned.pk, 'Assigned Teacher')],
        )
        self.assertEqual(payload['section_instructors'][0]['email'], primary.email)
        self.assertEqual(payload['section_instructors'][0]['contact_number'], '555-0101')
        self.assertEqual(payload['section_instructors'][0]['role'], User.RoleChoices.INSTRUCTOR)

    def test_student_can_update_editable_profile_fields(self):
        self.client.force_authenticate(user=self.student)

        response = self.client.patch('/api/users/update_profile/', {
            'email': 'updated@example.com',
            'contact_number': '09171234567',
        })

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()['email'], 'updated@example.com')
        self.assertEqual(response.json()['contact_number'], '09171234567')


class InstructorProfileTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.instructor = User.objects.create_user(
            username='instructor5',
            email='instructor5@example.com',
            password='secret123',
            role='instructor',
            first_name='Instructor',
            last_name='Five',
            instructor_id='I005',
            nfc_uid='NFC-I005',
        )

    def test_instructor_can_update_editable_profile_fields(self):
        self.client.force_authenticate(user=self.instructor)

        response = self.client.patch('/api/users/update_profile/', {
            'first_name': 'Updated',
            'last_name': 'Instructor',
            'instructor_id': 'I005-UPDATED',
            'username': 'updated-instructor5',
            'email': 'updated-instructor5@example.com',
            'contact_number': '09171234567',
            'role': 'admin',
        })

        self.assertEqual(response.status_code, 200)
        self.instructor.refresh_from_db()
        self.assertEqual(self.instructor.first_name, 'Updated')
        self.assertEqual(self.instructor.last_name, 'Instructor')
        self.assertEqual(self.instructor.instructor_id, 'I005-UPDATED')
        self.assertEqual(self.instructor.username, 'updated-instructor5')
        self.assertEqual(self.instructor.email, 'updated-instructor5@example.com')
        self.assertEqual(self.instructor.contact_number, '09171234567')
        self.assertEqual(self.instructor.role, 'instructor')

    def test_instructor_profile_includes_assigned_section_academic_details(self):
        section = Section.objects.create(
            section_name='Instructor Profile Section',
            subject_code='IP-301',
            program='Information Technology',
            year_level='3',
            academic_year='2026-2027',
            instructor=self.instructor,
        )
        self.instructor.assigned_sections.add(section)
        self.client.force_authenticate(user=self.instructor)

        response = self.client.get('/api/users/profile/')

        self.assertEqual(response.status_code, 200, response.json())
        self.assertEqual(response.json()['assigned_sections'], [{
            'section_id': section.pk,
            'section_name': 'Instructor Profile Section',
            'subject_code': 'IP-301',
            'program': 'Information Technology',
            'year_level': '3',
            'academic_year': '2026-2027',
        }])


class ProfileAdminFormTests(TestCase):
    def test_instructor_creation_form_does_not_require_student_id(self):
        from .admin import InstructorCreationForm

        form = InstructorCreationForm(data={
            'first_name': 'Test',
            'last_name': 'Instructor',
            'instructor_id': 'I100',
            'email': 'test-instructor@example.com',
            'nfc_uid': 'NFC-TEST-001',
            'password1': 'secret1234',
            'password2': 'secret1234',
            'is_active': True,
        })

        self.assertTrue(form.is_valid(), msg=form.errors)
        user = form.save()
        self.assertEqual(user.role, 'instructor')
        self.assertEqual(user.instructor_id, 'I100')
        self.assertIsNone(user.student_id)

    def test_student_creation_form_does_not_require_instructor_id(self):
        from .admin import StudentCreationForm

        form = StudentCreationForm(data={
            'first_name': 'Test',
            'last_name': 'Student',
            'student_id': 'S100',
            'email': 'test-student@example.com',
            'nfc_uid': 'NFC-STUD-001',
            'password1': 'secret1234',
            'password2': 'secret1234',
            'is_active': True,
        })

        self.assertTrue(form.is_valid(), msg=form.errors)
        user = form.save()
        self.assertEqual(user.role, 'student')
        self.assertEqual(user.student_id, 'S100')
        self.assertIsNone(user.instructor_id)

    def test_pre_enrollment_student_can_be_created_without_nfc_uid_or_password(self):
        from .admin import StudentCreationForm

        form = StudentCreationForm(data={
            'first_name': 'Roster',
            'last_name': 'Student',
            'student_id': 'ROSTER-100',
            'email': 'roster.student@example.com',
            'username': '',
            'nfc_uid': '',
            'password1': '',
            'password2': '',
            'is_active': True,
        })

        self.assertTrue(form.is_valid(), msg=form.errors)
        user = form.save()
        self.assertEqual(user.username, 'ROSTER-100')
        self.assertEqual(user.student_id, 'ROSTER-100')
        self.assertIsNone(user.nfc_uid)
        self.assertFalse(user.has_usable_password())

    def test_student_nfc_uid_remains_unique_when_present(self):
        User.objects.create_user(
            username='uid-owner',
            email='uid-owner@example.com',
            password=None,
            student_id='UID-OWNER-001',
            nfc_uid='NFC-UNIQUE-STUDENT-001',
        )
        duplicate = User(
            username='uid-duplicate',
            email='uid-duplicate@example.com',
            student_id='UID-DUPLICATE-001',
            nfc_uid='NFC-UNIQUE-STUDENT-001',
            role=User.RoleChoices.STUDENT,
        )
        duplicate.set_unusable_password()

        with self.assertRaises(ValidationError):
            duplicate.full_clean()

    def test_student_id_remains_required_without_nfc_uid(self):
        student = User(
            username='missing-student-id',
            email='missing-student-id@example.com',
            student_id='',
            nfc_uid=None,
            role=User.RoleChoices.STUDENT,
        )
        student.set_unusable_password()

        with self.assertRaises(ValidationError) as raised:
            student.full_clean()
        self.assertIn('student_id', raised.exception.message_dict)

    def test_admin_creation_form_does_not_create_mixed_role_data(self):
        from .admin import AdminCreationForm

        form = AdminCreationForm(data={
            'first_name': 'Test',
            'last_name': 'Admin',
            'email': 'test-admin@example.com',
            'nfc_uid': 'NFC-ADMIN-001',
            'password1': 'secret1234',
            'password2': 'secret1234',
            'is_active': True,
        })

        self.assertTrue(form.is_valid(), msg=form.errors)
        user = form.save()
        self.assertEqual(user.role, 'admin')
        self.assertIsNone(user.student_id)
        self.assertIsNone(user.instructor_id)
        self.assertFalse(user.assigned_sections.exists())
        from .models import AdminProfile
        self.assertIsInstance(user, AdminProfile)
        self.assertEqual(user._meta.model_name, 'adminprofile')

    def test_create_superuser_sets_admin_role_and_superuser_flags(self):
        from .models import User, AdminProfile, StudentProfile, InstructorProfile

        admin = User.objects.create_superuser(
            username='superadmin',
            email='superadmin@example.com',
            password='secret1234',
        )

        self.assertEqual(admin.role, User.RoleChoices.ADMIN)
        self.assertTrue(admin.is_staff)
        self.assertTrue(admin.is_superuser)
        self.assertFalse(StudentProfile.objects.filter(pk=admin.pk).exists())
        self.assertFalse(InstructorProfile.objects.filter(pk=admin.pk).exists())
        self.assertTrue(AdminProfile.objects.filter(pk=admin.pk).exists())
        self.assertEqual(admin._meta.model_name, 'user')
        self.assertEqual(admin._meta.app_label, 'api')

    def test_student_creation_form_saves_studentprofile_instance(self):
        from .admin import StudentCreationForm
        from .models import StudentProfile

        form = StudentCreationForm(data={
            'first_name': 'Test',
            'last_name': 'Student',
            'student_id': 'S200',
            'email': 'test-student2@example.com',
            'nfc_uid': 'NFC-STUD-002',
            'password1': 'secret1234',
            'password2': 'secret1234',
            'is_active': True,
        })

        self.assertTrue(form.is_valid(), msg=form.errors)
        user = form.save()
        self.assertEqual(user.role, 'student')
        self.assertIsInstance(user, StudentProfile)
        self.assertEqual(user._meta.model_name, 'studentprofile')

    def test_instructor_creation_form_saves_instructorprofile_instance(self):
        from .admin import InstructorCreationForm
        from .models import InstructorProfile

        form = InstructorCreationForm(data={
            'first_name': 'Test',
            'last_name': 'Instructor',
            'instructor_id': 'I200',
            'email': 'test-instructor2@example.com',
            'nfc_uid': 'NFC-INST-002',
            'password1': 'secret1234',
            'password2': 'secret1234',
            'is_active': True,
        })

        self.assertTrue(form.is_valid(), msg=form.errors)
        user = form.save()
        self.assertEqual(user.role, 'instructor')
        self.assertIsInstance(user, InstructorProfile)
        self.assertEqual(user._meta.model_name, 'instructorprofile')

    def test_instructor_admin_uses_autocomplete_for_assigned_sections(self):
        from .admin import InstructorProfileAdmin, SectionAdmin

        self.assertIn('assigned_sections', InstructorProfileAdmin.autocomplete_fields)
        self.assertIn('section_name', SectionAdmin.search_fields)
        self.assertIn('subject_code', SectionAdmin.search_fields)

    def test_section_admin_instructor_field_filters_instructors(self):
        from django.contrib import admin
        from .admin import SectionAdmin
        from .models import Section, User

        admin_instance = SectionAdmin(Section, admin.site)
        field = admin_instance.formfield_for_foreignkey(Section._meta.get_field('instructor'), None)
        self.assertEqual(field.queryset.filter(role=User.RoleChoices.INSTRUCTOR).count(), field.queryset.count())

    def test_activity_admin_instructor_fields_filter_to_instructors(self):
        from django.contrib import admin
        from .admin import ActivityAdmin
        from .models import Activity, User

        admin_instance = ActivityAdmin(Activity, admin.site)
        assigned_field = admin_instance.formfield_for_foreignkey(Activity._meta.get_field('assigned_instructor'), None)
        created_field = admin_instance.formfield_for_foreignkey(Activity._meta.get_field('created_by'), None)
        self.assertEqual(assigned_field.queryset.filter(role=User.RoleChoices.INSTRUCTOR).count(), assigned_field.queryset.count())
        self.assertEqual(created_field.queryset.filter(role=User.RoleChoices.INSTRUCTOR).count(), created_field.queryset.count())

    def test_submission_admin_student_field_filters_students(self):
        from django.contrib import admin
        from .admin import SubmissionAdmin
        from .models import Submission, User

        admin_instance = SubmissionAdmin(Submission, admin.site)
        field = admin_instance.formfield_for_foreignkey(Submission._meta.get_field('student'), None)
        self.assertEqual(field.queryset.filter(role=User.RoleChoices.STUDENT).count(), field.queryset.count())


class StudentWorkbookImportTests(TestCase):
    def setUp(self):
        self.client = Client()
        self.admin = User.objects.create_superuser(
            username='student-import-admin',
            email='student-import-admin@example.com',
            password='TestAdminPassword-123',
            nfc_uid='STUDENT-IMPORT-ADMIN-UID',
        )
        self.client.force_login(self.admin)
        self.section = Section.objects.create(subject_code='XLSX-TEST', section_name='Excel Import Test Section')
        self.import_url = reverse('admin:api_studentprofile_import_students')

    @staticmethod
    def workbook_file(headers, rows, filename='students.xlsx'):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(headers)
        for row in rows:
            worksheet.append(row)
        contents = BytesIO()
        workbook.save(contents)
        return SimpleUploadedFile(
            filename,
            contents.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )

    def upload(self, headers, rows):
        return self.client.post(
            self.import_url,
            {'workbook': self.workbook_file(headers, rows)},
        )

    def test_required_columns_are_validated(self):
        response = self.upload(['first_name', 'email'], [['Avery', 'avery@example.com']])

        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.context['errors'])
        self.assertFalse(User.objects.filter(email='avery@example.com').exists())

    def test_students_admin_changelist_links_to_import_page(self):
        response = self.client.get(reverse('admin:api_studentprofile_changelist'))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Import Students (.xlsx)')
        self.assertContains(response, self.import_url)

    def test_downloadable_template_uses_supported_columns_and_fake_data(self):
        from pathlib import Path
        from .student_import import read_student_workbook, validate_student_rows
        from .serializers import SectionSerializer

        template_path = Path(__file__).parent / 'static' / 'api' / 'student_import_template.xlsx'
        self.assertTrue(template_path.exists())
        template_workbook = load_workbook(template_path, read_only=True, data_only=True)
        self.assertEqual(
            next(template_workbook.active.iter_rows(values_only=True)),
            ('student_id', 'first_name', 'last_name', 'email', 'username', 'subject_code', 'role', 'active'),
        )
        template_workbook.close()
        uploaded = SimpleUploadedFile(template_path.name, template_path.read_bytes())
        rows, errors = read_student_workbook(uploaded)
        prepared, row_errors = validate_student_rows(rows)

        self.assertFalse(errors)
        self.assertFalse(row_errors)
        self.assertEqual(len(prepared), 1)
        self.assertEqual(prepared[0]['student_id'], 'DEMO-STUDENT-001')
        self.assertEqual(prepared[0]['email'], 'avery.example@example.invalid')
        self.assertEqual(SectionSerializer(self.section).data['subject_code'], 'XLSX-TEST')
        self.assertEqual(Section._meta.get_field('subject_code').verbose_name, 'Subject Code')
        self.assertFalse(User.objects.filter(student_id='DEMO-STUDENT-001').exists())

    def test_duplicate_student_id_and_email_inside_workbook_are_rejected(self):
        headers = ['student_id', 'first_name', 'last_name', 'email']
        rows = [
            ['DUP-100', 'Avery', 'Example', 'same@example.com'],
            ['DUP-100', 'Casey', 'Example', 'same@example.com'],
        ]
        response = self.upload(headers, rows)

        self.assertTrue(response.context['errors'])
        self.assertTrue(any(error['row'] == 3 and error['column'] == 'student_id' for error in response.context['errors']))
        self.assertTrue(any(error['row'] == 3 and error['column'] == 'email' for error in response.context['errors']))
        self.assertFalse(User.objects.filter(student_id='DUP-100').exists())
        self.assertFalse(User.objects.filter(email='same@example.com').exists())

    def test_existing_student_id_is_rejected_without_overwriting(self):
        existing = User.objects.create_user(
            username='existing-student',
            email='existing.student@example.com',
            password=None,
            student_id='EXISTING-100',
            role=User.RoleChoices.STUDENT,
        )
        response = self.upload(
            ['student_id', 'first_name', 'last_name', 'email'],
            [['EXISTING-100', 'Changed', 'Name', 'changed@example.com']],
        )

        self.assertTrue(response.context['errors'])
        existing.refresh_from_db()
        self.assertEqual(existing.first_name, '')
        self.assertFalse(User.objects.filter(email='changed@example.com').exists())

    def test_existing_email_is_rejected(self):
        User.objects.create_user(
            username='existing-email-student',
            email='already.used@example.com',
            password=None,
            student_id='EMAIL-100',
            role=User.RoleChoices.STUDENT,
        )
        response = self.upload(
            ['student_id', 'first_name', 'last_name', 'email'],
            [['EMAIL-101', 'Avery', 'Example', 'already.used@example.com']],
        )

        self.assertTrue(any(error['column'] == 'email' for error in response.context['errors']))
        self.assertFalse(User.objects.filter(student_id='EMAIL-101').exists())

    def test_existing_username_is_rejected(self):
        User.objects.create_user(
            username='chosen_username',
            email='username.owner@example.com',
            password=None,
            student_id='USERNAME-OWNER-100',
            role=User.RoleChoices.STUDENT,
        )
        response = self.upload(
            ['student_id', 'first_name', 'last_name', 'email', 'username'],
            [['USERNAME-NEW-100', 'Avery', 'Example', 'new.username@example.com', 'chosen_username']],
        )

        self.assertTrue(any(error['column'] == 'username' for error in response.context['errors']))
        self.assertFalse(User.objects.filter(student_id='USERNAME-NEW-100').exists())

    def test_invalid_subject_code_reports_row_and_imports_nothing(self):
        response = self.upload(
            ['student_id', 'first_name', 'last_name', 'email', 'subject_code'],
            [
                ['SECTION-VALID-100', 'Casey', 'Example', 'casey.section@example.com', ''],
                ['SECTION-INVALID-100', 'Avery', 'Example', 'avery.section@example.com', 'MISSING-SECTION'],
            ],
        )

        self.assertTrue(any(error['row'] == 3 and error['column'] == 'subject_code' for error in response.context['errors']))
        self.assertFalse(User.objects.filter(student_id__in=['SECTION-VALID-100', 'SECTION-INVALID-100']).exists())

    def test_ambiguous_subject_code_is_rejected(self):
        Section.objects.create(subject_code='DUPLICATE-CODE', section_name='Duplicate Section One')
        Section.objects.create(subject_code='duplicate-code', section_name='Duplicate Section Two')
        response = self.upload(
            ['student_id', 'first_name', 'last_name', 'email', 'subject_code'],
            [['SECTION-DUP-100', 'Avery', 'Example', 'section.duplicate@example.com', 'Duplicate-Code']],
        )

        self.assertTrue(any(error['column'] == 'subject_code' and 'More than one' in error['message'] for error in response.context['errors']))
        self.assertFalse(User.objects.filter(student_id='SECTION-DUP-100').exists())

    def test_nfc_uid_and_password_columns_are_not_accepted(self):
        response = self.upload(
            ['student_id', 'first_name', 'last_name', 'email', 'nfc_uid', 'password'],
            [['SECURITY-100', 'Avery', 'Example', 'security@example.com', 'NFC-UID', 'do-not-import']],
        )

        self.assertTrue(any(error['column'] == 'header' and 'Unsupported column' in error['message'] for error in response.context['errors']))
        self.assertFalse(User.objects.filter(student_id='SECURITY-100').exists())

    def test_invalid_email_and_active_values_are_rejected(self):
        response = self.upload(
            ['student_id', 'first_name', 'last_name', 'email', 'active'],
            [['BAD-100', 'Avery', 'Example', 'not-an-email', 'sometimes']],
        )

        self.assertEqual({error['column'] for error in response.context['errors']}, {'email', 'active'})
        self.assertFalse(User.objects.filter(student_id='BAD-100').exists())

    def test_duplicate_generated_username_is_rejected(self):
        response = self.upload(
            ['student_id', 'first_name', 'last_name', 'email'],
            [
                ['SAME USER', 'Avery', 'Example', 'avery.one@example.com'],
                ['SAME_USER', 'Casey', 'Example', 'casey.two@example.com'],
            ],
        )

        self.assertTrue(any(error['column'] == 'username' and error['row'] == 3 for error in response.context['errors']))
        self.assertFalse(User.objects.filter(student_id__in=['SAME USER', 'SAME_USER']).exists())

    def test_successful_import_is_confirmed_atomically_and_uses_unusable_passwords(self):
        response = self.upload(
            ['student_id', 'first_name', 'last_name', 'email', 'username', 'subject_code', 'role', 'active'],
            [
                ['IMPORT-100', 'Avery', 'Example', 'avery.import@example.com', '', 'XLSX-TEST', 'student', 'true'],
                ['IMPORT-101', 'Casey', 'Example', 'casey.import@example.com', 'casey_import_user', '', 'student', 'false'],
            ],
        )
        self.assertFalse(response.context['errors'])
        self.assertTrue(response.context['confirm_nonce'])
        self.assertEqual(User.objects.filter(student_id__in=['IMPORT-100', 'IMPORT-101']).count(), 0)

        confirmed = self.client.post(self.import_url, {
            'action': 'confirm',
            'confirm_nonce': response.context['confirm_nonce'],
        })

        self.assertEqual(confirmed.context['result'], {'created': 2, 'updated': 0})
        students = list(User.objects.filter(student_id__in=['IMPORT-100', 'IMPORT-101']).order_by('student_id'))
        self.assertEqual(len(students), 2)
        self.assertEqual(students[0].username, 'student_IMPORT-100')
        self.assertEqual(students[0].section, self.section)
        self.assertTrue(students[0].is_active)
        self.assertIsNone(students[0].nfc_uid)
        self.assertFalse(students[0].has_usable_password())
        self.assertFalse(students[1].is_active)
        self.assertIsNone(students[1].section)

    @override_settings(TAPTRACK_NFC_DEVICE_MAP=[
        {'device_id': 'CABINET1-STATION1', 'api_key': 'unit-test-device-key', 'station': 'Station 1', 'cabinet_name': 'Cabinet 1', 'active': True},
    ])
    def test_imported_student_can_complete_existing_nfc_enrollment_flow(self):
        upload = self.upload(
            ['student_id', 'first_name', 'last_name', 'email'],
            [['ENROLL-100', 'Taylor', 'Example', 'taylor.enroll@example.com']],
        )
        self.client.post(self.import_url, {
            'action': 'confirm',
            'confirm_nonce': upload.context['confirm_nonce'],
        })
        student = User.objects.get(student_id='ENROLL-100')
        test_uid = 'EXCEL-ENROLL-UID-100'
        device_key = 'unit-test-device-key'
        previous_key = os.environ.get('DEVICE_API_KEY')
        os.environ['DEVICE_API_KEY'] = device_key
        try:
            initial = self.client.post('/api/verify-nfc/', {'nfc_uid': test_uid}, HTTP_X_API_KEY=device_key)
            self.assertEqual(initial.status_code, 404, initial.json())
            token = parse_qs(urlparse(initial.json()['registration_url']).query)['token'][0]
            registration = self.client.post('/api/cabinet/enrollment/register/', {
                'token': token,
                'student_id': student.student_id,
                'full_name': 'Taylor Example',
                'email': student.email,
                'password': 'Enrollment-Test-Password-123',
                'confirm_password': 'Enrollment-Test-Password-123',
            })
            self.assertEqual(registration.status_code, 201, registration.json())
            registration_log = AccessLog.objects.get(
                nfc_uid=test_uid,
                action=AccessLog.AccessActionChoices.REGISTRATION,
            )
            self.assertEqual(registration_log.user, student)
            self.assertEqual(registration_log.status, AccessLog.AccessStatusChoices.SUCCESS)
            self.assertEqual(registration_log.station, 'Station 1')
            self.assertEqual(registration_log.cabinet_name, 'Cabinet 1')
            self.assertEqual(registration_log.reason, 'NFC card registered successfully')
            verification = self.client.post('/api/verify-nfc/', {'nfc_uid': test_uid}, HTTP_X_API_KEY=device_key)
            self.assertEqual(verification.status_code, 200, verification.json())
            self.assertEqual(verification.json()['student_id'], 'ENROLL-100')
            self.assertEqual(verification.json()['role'], 'student')
        finally:
            if previous_key is None:
                os.environ.pop('DEVICE_API_KEY', None)
            else:
                os.environ['DEVICE_API_KEY'] = previous_key


class InstructorAccessLogTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.instructor = User.objects.create_user(
            username='access-log-instructor',
            email='access-log-instructor@example.com',
            password='secret1234',
            role=User.RoleChoices.INSTRUCTOR,
            instructor_id='ACCESS-LOG-INSTRUCTOR',
        )
        self.section_a = Section.objects.create(section_name='Log Section A', instructor=self.instructor)
        self.section_b = Section.objects.create(section_name='Log Section B')
        self.section_b.assigned_instructors.add(self.instructor)
        self.unrelated_section = Section.objects.create(section_name='Unrelated Log Section')
        self.student_a = User.objects.create_user(
            username='access-log-student-a',
            email='access-log-student-a@example.com',
            password='secret1234',
            role=User.RoleChoices.STUDENT,
            student_id='LOG-STUDENT-A',
            first_name='Casey',
            last_name='Alpha',
            section=self.section_a,
        )
        self.student_b = User.objects.create_user(
            username='access-log-student-b',
            email='access-log-student-b@example.com',
            password='secret1234',
            role=User.RoleChoices.STUDENT,
            student_id='LOG-STUDENT-B',
            first_name='Jordan',
            last_name='Beta',
            section=self.section_b,
        )
        self.outside_student = User.objects.create_user(
            username='access-log-outside-student',
            email='access-log-outside-student@example.com',
            password='secret1234',
            role=User.RoleChoices.STUDENT,
            student_id='LOG-STUDENT-OUTSIDE',
            section=self.unrelated_section,
        )
        now = timezone.now()
        self.success_log = self.create_log(
            self.student_a,
            AccessLog.AccessStatusChoices.SUCCESS,
            nfc_uid='AA11BB22',
            station='Reader 1',
            cabinet_name='Cabinet Alpha',
            reason='NFC verified successfully',
            access_time=now,
        )
        self.rejected_log = self.create_log(
            self.student_b,
            AccessLog.AccessStatusChoices.REJECTED,
            nfc_uid='',
            station='Reader 2',
            cabinet_name='',
            reason='',
            access_time=now,
        )
        self.duplicate_log = self.create_log(
            self.student_b,
            AccessLog.AccessStatusChoices.DUPLICATE,
            nfc_uid='CC33DD44',
            station='Reader 2',
            cabinet_name='Cabinet Gamma',
            reason='Duplicate participant scan',
            access_time=now,
        )
        self.old_log = self.create_log(
            self.student_a,
            AccessLog.AccessStatusChoices.SUCCESS,
            nfc_uid='EE55FF66',
            station='Reader 1',
            cabinet_name='Cabinet Old',
            reason='NFC verified successfully',
            access_time=now - timedelta(days=2),
        )
        self.outside_log = self.create_log(
            self.outside_student,
            AccessLog.AccessStatusChoices.SUCCESS,
            nfc_uid='OUTSIDE-UID',
            station='Reader Outside',
            cabinet_name='Cabinet Outside',
            reason='NFC verified successfully',
            access_time=now,
        )
        AccessLogPhoto.objects.create(
            access_log=self.success_log,
            capture_status=AccessLogPhoto.CaptureStatusChoices.SUCCESS,
            captured_at=now,
        )
        AccessLogPhoto.objects.create(
            access_log=self.duplicate_log,
            capture_status=AccessLogPhoto.CaptureStatusChoices.FAILED,
            diagnostic_error='Camera unavailable',
        )

    def create_log(self, user, status, **fields):
        access_time = fields.pop('access_time')
        log = AccessLog.objects.create(user=user, status=status, **fields)
        log.access_time = access_time
        log.save(update_fields=['access_time'])
        return log

    def setUpInstructor(self):
        self.client.force_authenticate(user=self.instructor)

    def test_instructor_access_logs_serialize_actual_record_and_limit_scope(self):
        self.setUpInstructor()

        response = self.client.get('/api/access-logs/?page_size=25')

        self.assertEqual(response.status_code, 200, response.json())
        results = response.json()['results']
        self.assertEqual({entry['id'] for entry in results}, {
            self.success_log.pk,
            self.rejected_log.pk,
            self.duplicate_log.pk,
            self.old_log.pk,
        })
        serialized = next(entry for entry in results if entry['id'] == self.success_log.pk)
        self.assertEqual(serialized['student_id'], 'LOG-STUDENT-A')
        self.assertEqual(serialized['student_name'], 'Casey Alpha')
        self.assertEqual(serialized['section_name'], 'Log Section A')
        self.assertEqual(serialized['nfc_uid'], 'AA11BB22')
        self.assertEqual(serialized['cabinet_name'], 'Cabinet Alpha')
        self.assertEqual(serialized['station'], 'Reader 1')
        self.assertEqual(serialized['access_result'], 'Success')
        self.assertEqual(serialized['photo_capture_status'], 'success')
        self.assertNotIn('image_endpoint', serialized)

        failed = next(entry for entry in results if entry['id'] == self.rejected_log.pk)
        self.assertEqual(failed['access_result'], 'Failed')
        self.assertEqual(failed['nfc_uid'], '')
        self.assertEqual(failed['cabinet_name'], '')
        self.assertEqual(failed['reason'], '')
        self.assertEqual(next(entry for entry in results if entry['id'] == self.duplicate_log.pk)['reason'], 'Duplicate participant scan')

    def test_instructor_filters_cover_failure_groups_search_dates_sections_and_pages(self):
        self.setUpInstructor()
        today = timezone.localdate().isoformat()

        failed_response = self.client.get('/api/access-logs/?status=failed&page_size=25')
        self.assertEqual(failed_response.status_code, 200, failed_response.json())
        self.assertEqual(
            {entry['id'] for entry in failed_response.json()['results']},
            {self.rejected_log.pk, self.duplicate_log.pk},
        )
        success_response = self.client.get(f'/api/access-logs/?status=success&access_time_after={today}&access_time_before={today}')
        self.assertEqual([entry['id'] for entry in success_response.json()['results']], [self.success_log.pk])

        cabinet_search = self.client.get('/api/access-logs/?search=Gamma')
        self.assertEqual([entry['id'] for entry in cabinet_search.json()['results']], [self.duplicate_log.pk])
        name_search = self.client.get('/api/access-logs/?search=Casey Alpha')
        self.assertEqual(
            {entry['id'] for entry in name_search.json()['results']},
            {self.success_log.pk, self.old_log.pk},
        )
        nfc_search = self.client.get('/api/access-logs/?search=AA11BB22')
        self.assertEqual([entry['id'] for entry in nfc_search.json()['results']], [self.success_log.pk])
        section_response = self.client.get(f'/api/access-logs/?user__section={self.section_b.pk}')
        self.assertEqual(
            {entry['id'] for entry in section_response.json()['results']},
            {self.rejected_log.pk, self.duplicate_log.pk},
        )
        photo_response = self.client.get('/api/access-logs/?photo_capture_status=failed')
        self.assertEqual(photo_response.status_code, 200, photo_response.json())
        self.assertEqual([entry['id'] for entry in photo_response.json()['results']], [self.duplicate_log.pk])
        date_response = self.client.get(f'/api/access-logs/?access_time_after={today}&access_time_before={today}')
        self.assertEqual(
            {entry['id'] for entry in date_response.json()['results']},
            {self.success_log.pk, self.rejected_log.pk, self.duplicate_log.pk},
        )
        one_item_page = self.client.get('/api/access-logs/?page_size=1&page=2')
        self.assertEqual(one_item_page.json()['count'], 4)
        self.assertEqual(len(one_item_page.json()['results']), 1)

    def test_cabinet_filter_options_are_instructor_and_section_scoped(self):
        self.setUpInstructor()

        all_options = self.client.get('/api/access-logs/filter-options/')
        self.assertEqual(all_options.status_code, 200, all_options.json())
        self.assertEqual(set(all_options.json()['cabinets']), {'Cabinet Alpha', 'Cabinet Gamma', 'Cabinet Old'})

        section_options = self.client.get(f'/api/access-logs/filter-options/?user__section={self.section_b.pk}')
        self.assertEqual(section_options.json()['cabinets'], ['Cabinet Gamma'])


class StudentAccessLogsTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.student = User.objects.create_user(
            username='student1',
            email='student1@example.com',
            password='secret123',
            role='student',
            first_name='Student',
            last_name='One',
            student_id='S001',
        )
        self.other_student = User.objects.create_user(
            username='student2',
            email='student2@example.com',
            password='secret123',
            role='student',
            first_name='Student',
            last_name='Two',
            student_id='S002',
        )
        AccessLog.objects.create(user=self.student, status='success', cabinet_name='Cabinet A', reason='Entry granted', nfc_uid='ABC123')
        AccessLog.objects.create(user=self.other_student, status='failed', cabinet_name='Cabinet B', reason='Denied', nfc_uid='XYZ999')

    def test_student_only_sees_own_access_logs(self):
        self.client.force_authenticate(user=self.student)

        response = self.client.get('/api/access-logs/')

        self.assertEqual(response.status_code, 200)
        data = response.json()
        results = data.get('results', [])
        self.assertEqual(len(results), 1)
        self.assertEqual(results[0]['cabinet_name'], 'Cabinet A')
        self.assertEqual(results[0]['user'], self.student.id)

    def test_student_stats_endpoint_returns_private_totals(self):
        self.client.force_authenticate(user=self.student)

        response = self.client.get('/api/access-logs/stats/')

        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data['total_accesses_today'], 1)
        self.assertEqual(data['successful_accesses_today'], 1)
        self.assertEqual(data['failed_accesses_today'], 0)

    def test_access_logs_store_nfc_uid_and_cabinet_session(self):
        session = CabinetSession.objects.create(station='Station 1', status='open')
        session.opened_by.add(self.student)
        log = AccessLog.objects.create(
            user=self.student,
            nfc_uid='ABCD-1234',
            station='Station 1',
            action='open',
            status='success',
            reason='Cabinet session opened',
            cabinet_session=session,
        )

        self.assertEqual(log.nfc_uid, 'ABCD-1234')
        self.assertEqual(log.station, 'Station 1')
        self.assertEqual(log.cabinet_session_id, session.id)

    def test_access_log_api_rejects_client_created_audit_values(self):
        admin_user = User.objects.create_superuser(
            username='access-log-api-admin',
            email='access-log-api-admin@example.com',
            password='secret1234',
        )
        self.client.force_authenticate(user=admin_user)

        response = self.client.post('/api/access-logs/', {
            'user': self.student.pk,
            'status': 'success',
            'action': 'open',
            'nfc_uid': 'CLIENT-SUPPLIED',
            'station': 'Station 1',
            'cabinet_name': 'Cabinet 1',
            'reason': 'Forged audit data',
        }, format='json')

        self.assertEqual(response.status_code, 405)
        self.assertFalse(AccessLog.objects.filter(nfc_uid='CLIENT-SUPPLIED').exists())

    def test_access_log_admin_is_read_only(self):
        access_log_admin = AccessLogAdmin(model=AccessLog, admin_site=django_admin.site)
        admin_user = User.objects.create_superuser(
            username='access-log-admin',
            email='access-log-admin@example.com',
            password='secret1234',
        )
        request = RequestFactory().get('/admin/api/accesslog/')
        request.user = admin_user

        self.assertFalse(access_log_admin.has_add_permission(request))
        self.assertFalse(access_log_admin.has_change_permission(request))
        self.assertFalse(access_log_admin.has_delete_permission(request))
        self.assertTrue(access_log_admin.has_view_permission(request))
        self.assertEqual(access_log_admin.get_model_perms(request), {'view': True})
        self.assertTrue({
            'user', 'cabinet_session', 'access_time', 'status', 'action',
            'nfc_uid', 'station', 'cabinet_name', 'reason', 'updated_at',
        }.issubset(set(access_log_admin.get_readonly_fields(request))))

        session = CabinetSession.objects.create(station='Station 1', status=CabinetSession.StatusChoices.OPEN)
        log = AccessLog.objects.create(
            user=self.student,
            cabinet_session=session,
            access_time=timezone.now(),
            status='success',
            action='scan',
            nfc_uid='ABCD1234',
            station='Station 1',
            cabinet_name='Cabinet 1',
            reason='Original audit reason',
        )
        self.client.force_login(admin_user)
        response = self.client.get(reverse('admin:api_accesslog_change', args=[log.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'ABCD1234')
        self.assertContains(response, 'Original audit reason')

        response = self.client.post(reverse('admin:api_accesslog_change', args=[log.pk]), {
            'user': self.other_student.pk,
            'cabinet_session': '',
            'access_time': '2000-01-01 00:00:00',
            'status': 'failed',
            'action': 'open',
            'nfc_uid': 'OVERRIDE',
            'station': 'Station 2',
            'cabinet_name': 'Forged cabinet',
            'reason': 'Forged reason',
            'updated_at': '2000-01-01 00:00:00',
        })
        self.assertEqual(response.status_code, 403)
        log.refresh_from_db()
        self.assertEqual(log.user, self.student)
        self.assertEqual(log.cabinet_session, session)
        self.assertEqual(log.status, 'success')
        self.assertEqual(log.action, 'scan')
        self.assertEqual(log.nfc_uid, 'ABCD1234')
        self.assertEqual(log.station, 'Station 1')
        self.assertEqual(log.cabinet_name, 'Cabinet 1')
        self.assertEqual(log.reason, 'Original audit reason')


@override_settings(TAPTRACK_NFC_DEVICE_MAP=[
    {'device_id': 'CABINET1-STATION1', 'api_key': 'test-device-key', 'station': 'Station 1', 'cabinet_name': 'Cabinet 1', 'active': True},
])
class PasswordResetFlowTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.device_api_key = 'test-device-key'
        os.environ['DEVICE_API_KEY'] = self.device_api_key

        self.student_a = User.objects.create_user(
            username='student-a',
            email='student-a@example.com',
            password='old-password-1',
            role=User.RoleChoices.STUDENT,
            first_name='Student',
            last_name='A',
            student_id='STU-001',
            nfc_uid='NFC-USER-001',
        )
        self.student_b = User.objects.create_user(
            username='student-b',
            email='student-b@example.com',
            password='old-password-2',
            role=User.RoleChoices.STUDENT,
            first_name='Student',
            last_name='B',
            student_id='STU-002',
            nfc_uid='NFC-USER-002',
        )

    def test_password_reset_request_is_bound_to_existing_user_and_requires_nfc_verification(self):
        response = self.client.post('/api/auth/password-reset/request/', {
            'student_id': 'STU-001',
            'email': 'student-a@example.com',
        }, format='json')

        self.assertEqual(response.status_code, 200, response.json())
        payload = response.json()
        self.assertIn('request_id', payload)

        reset_request = PasswordResetRequest.objects.get(request_id=payload['request_id'])
        self.assertEqual(reset_request.user_id, self.student_a.pk)

        handoff_response = self.client.post('/api/auth/password-reset/cabinet-handoff/', {
            'request_id': reset_request.request_id,
        }, HTTP_X_API_KEY=self.device_api_key, format='json')
        self.assertEqual(handoff_response.status_code, 200, handoff_response.json())

        verify_response = self.client.post('/api/auth/password-reset/verify-nfc/', {
            'request_id': reset_request.request_id,
            'student_id': 'STU-001',
            'nfc_uid': 'NFC-USER-001',
        }, HTTP_X_API_KEY=self.device_api_key, format='json')

        self.assertEqual(verify_response.status_code, 200, verify_response.json())
        verify_payload = verify_response.json()
        self.assertIn('reset_token', verify_payload)

        validation_response = self.client.post('/api/auth/password-reset/validate-token/', {
            'request_id': reset_request.request_id,
            'reset_token': verify_payload['reset_token'],
        }, format='json')

        self.assertEqual(validation_response.status_code, 200, validation_response.json())
        validation_payload = validation_response.json()
        self.assertTrue(validation_payload['valid'])
        self.assertIn('reset_authorization', validation_payload)
        self.assertNotIn('reset_token', validation_payload)

        confirm_response = self.client.post('/api/auth/password-reset/confirm/', {
            'request_id': reset_request.request_id,
            'reset_authorization': validation_payload['reset_authorization'],
            'new_password': 'NewPassword123!',
            'confirm_password': 'NewPassword123!',
        }, format='json')

        self.assertEqual(confirm_response.status_code, 200, confirm_response.json())
        self.student_a.refresh_from_db()
        self.assertTrue(self.student_a.check_password('NewPassword123!'))

    def test_reset_request_rejects_mismatched_student_id_or_nfc_uid(self):
        response = self.client.post('/api/auth/password-reset/request/', {
            'student_id': 'STU-002',
            'email': 'student-b@example.com',
        }, format='json')

        self.assertEqual(response.status_code, 200, response.json())
        reset_request = PasswordResetRequest.objects.get(request_id=response.json()['request_id'])

        handoff_response = self.client.post('/api/auth/password-reset/cabinet-handoff/', {
            'request_id': reset_request.request_id,
        }, HTTP_X_API_KEY=self.device_api_key, format='json')
        self.assertEqual(handoff_response.status_code, 200, handoff_response.json())

        mismatched_response = self.client.post('/api/auth/password-reset/verify-nfc/', {
            'request_id': reset_request.request_id,
            'student_id': 'STU-001',
            'nfc_uid': 'NFC-USER-002',
        }, HTTP_X_API_KEY=self.device_api_key, format='json')

        self.assertEqual(mismatched_response.status_code, 403, mismatched_response.json())
        self.assertIn('does not match', mismatched_response.json()['error'])

    def _create_verified_reset_request(self, user=None):
        target = user or self.student_a
        request_response = self.client.post('/api/auth/password-reset/request/', {
            'student_id': target.student_id,
            'email': target.email,
        }, format='json')
        reset_request = PasswordResetRequest.objects.get(request_id=request_response.json()['request_id'])
        handoff_response = self.client.post('/api/auth/password-reset/cabinet-handoff/', {
            'request_id': reset_request.request_id,
        }, HTTP_X_API_KEY=self.device_api_key, format='json')
        if handoff_response.status_code != 200:
            raise AssertionError(handoff_response.json())
        verify_response = self.client.post('/api/auth/password-reset/verify-nfc/', {
            'request_id': reset_request.request_id,
            'student_id': target.student_id,
            'nfc_uid': target.nfc_uid,
        }, HTTP_X_API_KEY=self.device_api_key, format='json')
        return reset_request, verify_response.json()['reset_token']

    def test_invalid_token_is_rejected_without_authorization(self):
        reset_request, _ = self._create_verified_reset_request()

        response = self.client.post('/api/auth/password-reset/validate-token/', {
            'request_id': reset_request.request_id,
            'reset_token': 'invalid-token',
        }, format='json')

        self.assertEqual(response.status_code, 400)
        self.assertNotIn('reset_authorization', response.json())

    def test_expired_token_is_rejected(self):
        reset_request, reset_token = self._create_verified_reset_request()
        reset_request.expires_at = timezone.now() - timedelta(minutes=1)
        reset_request.save(update_fields=['expires_at'])

        response = self.client.post('/api/auth/password-reset/validate-token/', {
            'request_id': reset_request.request_id,
            'reset_token': reset_token,
        }, format='json')

        self.assertEqual(response.status_code, 410)

    def test_token_cannot_validate_against_another_request(self):
        first_request, reset_token = self._create_verified_reset_request(self.student_a)
        second_request, _ = self._create_verified_reset_request(self.student_b)

        response = self.client.post('/api/auth/password-reset/validate-token/', {
            'request_id': second_request.request_id,
            'reset_token': reset_token,
        }, format='json')

        self.assertEqual(response.status_code, 400)
        first_request.refresh_from_db()
        self.assertEqual(first_request.status, PasswordResetRequest.StatusChoices.VERIFIED)

    def test_used_token_is_rejected(self):
        reset_request, reset_token = self._create_verified_reset_request()
        reset_request.status = PasswordResetRequest.StatusChoices.USED
        reset_request.used_at = timezone.now()
        reset_request.save(update_fields=['status', 'used_at'])

        response = self.client.post('/api/auth/password-reset/validate-token/', {
            'request_id': reset_request.request_id,
            'reset_token': reset_token,
        }, format='json')

        self.assertEqual(response.status_code, 409)

    def test_expired_authorization_cannot_confirm_password(self):
        reset_request, reset_token = self._create_verified_reset_request()
        validation_response = self.client.post('/api/auth/password-reset/validate-token/', {
            'request_id': reset_request.request_id,
            'reset_token': reset_token,
        }, format='json')
        reset_request.refresh_from_db()
        reset_request.reset_authorization_expires_at = timezone.now() - timedelta(minutes=1)
        reset_request.save(update_fields=['reset_authorization_expires_at'])

        response = self.client.post('/api/auth/password-reset/confirm/', {
            'request_id': reset_request.request_id,
            'reset_authorization': validation_response.json()['reset_authorization'],
            'new_password': 'NewPassword123!',
            'confirm_password': 'NewPassword123!',
        }, format='json')

        self.assertEqual(response.status_code, 400)

    def test_reset_authorization_is_single_use(self):
        reset_request, reset_token = self._create_verified_reset_request()
        validation_response = self.client.post('/api/auth/password-reset/validate-token/', {
            'request_id': reset_request.request_id,
            'reset_token': reset_token,
        }, format='json')
        authorization = validation_response.json()['reset_authorization']

        response = self.client.post('/api/auth/password-reset/confirm/', {
            'request_id': reset_request.request_id,
            'reset_authorization': authorization,
            'new_password': 'NewPassword123!',
            'confirm_password': 'NewPassword123!',
        }, format='json')
        self.assertEqual(response.status_code, 200, response.json())

        replay_response = self.client.post('/api/auth/password-reset/confirm/', {
            'request_id': reset_request.request_id,
            'reset_authorization': authorization,
            'new_password': 'AnotherPassword123!',
            'confirm_password': 'AnotherPassword123!',
        }, format='json')
        self.assertEqual(replay_response.status_code, 400)

    def test_nfc_endpoint_requires_device_api_key(self):
        reset_request, _ = self._create_verified_reset_request()

        response = self.client.post('/api/auth/password-reset/verify-nfc/', {
            'request_id': reset_request.request_id,
            'student_id': self.student_a.student_id,
            'nfc_uid': self.student_a.nfc_uid,
        }, format='json')

        self.assertEqual(response.status_code, 403)

    def test_password_reset_verify_requires_claimed_cabinet_handoff(self):
        request_response = self.client.post('/api/auth/password-reset/request/', {
            'student_id': self.student_a.student_id,
            'email': self.student_a.email,
        }, format='json')
        request_id = request_response.json()['request_id']

        response = self.client.post('/api/auth/password-reset/verify-nfc/', {
            'request_id': request_id,
            'student_id': self.student_a.student_id,
            'nfc_uid': self.student_a.nfc_uid,
        }, HTTP_X_API_KEY=self.device_api_key, format='json')

        self.assertEqual(response.status_code, 409, response.json())

    def test_cabinet_handoff_requires_device_api_key_and_valid_request(self):
        request_response = self.client.post('/api/auth/password-reset/request/', {
            'student_id': self.student_a.student_id,
            'email': self.student_a.email,
        }, format='json')
        request_id = request_response.json()['request_id']

        missing_key_response = self.client.post('/api/auth/password-reset/cabinet-handoff/', {
            'request_id': request_id,
        }, format='json')
        self.assertEqual(missing_key_response.status_code, 403)

        invalid_response = self.client.post('/api/auth/password-reset/cabinet-handoff/', {
            'request_id': 'not-a-real-request',
        }, HTTP_X_API_KEY=self.device_api_key, format='json')
        self.assertEqual(invalid_response.status_code, 404)

    def test_cabinet_handoff_is_single_active_claim_and_expires(self):
        request_response = self.client.post('/api/auth/password-reset/request/', {
            'student_id': self.student_a.student_id,
            'email': self.student_a.email,
        }, format='json')
        reset_request = PasswordResetRequest.objects.get(request_id=request_response.json()['request_id'])

        first_response = self.client.post('/api/auth/password-reset/cabinet-handoff/', {
            'request_id': reset_request.request_id,
        }, HTTP_X_API_KEY=self.device_api_key, format='json')
        self.assertEqual(first_response.status_code, 200, first_response.json())

        second_response = self.client.post('/api/auth/password-reset/cabinet-handoff/', {
            'request_id': reset_request.request_id,
        }, HTTP_X_API_KEY=self.device_api_key, format='json')
        self.assertEqual(second_response.status_code, 409, second_response.json())

        reset_request.refresh_from_db()
        reset_request.cabinet_handoff_expires_at = timezone.now() - timedelta(seconds=1)
        reset_request.save(update_fields=['cabinet_handoff_expires_at'])
        retry_response = self.client.post('/api/auth/password-reset/cabinet-handoff/', {
            'request_id': reset_request.request_id,
        }, HTTP_X_API_KEY=self.device_api_key, format='json')
        self.assertEqual(retry_response.status_code, 200, retry_response.json())

    def test_normal_login_and_existing_nfc_access_still_work(self):
        login_response = self.client.post('/api/auth/login/', {
            'username': self.student_a.username,
            'password': 'old-password-1',
            'expected_role': 'student',
        }, format='json')
        self.assertEqual(login_response.status_code, 200, login_response.json())

        nfc_response = self.client.post('/api/verify-nfc/', {
            'nfc_uid': self.student_a.nfc_uid,
        }, HTTP_X_API_KEY=self.device_api_key, format='json')
        self.assertEqual(nfc_response.status_code, 200, nfc_response.json())

    def test_password_reset_request_still_works(self):
        response = self.client.post('/api/auth/password-reset/request/', {
            'student_id': self.student_b.student_id,
            'email': self.student_b.email,
        }, format='json')

        self.assertEqual(response.status_code, 200, response.json())
        self.assertIn('request_id', response.json())


class CabinetEventAPITests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.admin = User.objects.create_user(
            username='cabinet-event-admin',
            email='cabinet-event-admin@example.com',
            password='secret1234',
            role=User.RoleChoices.ADMIN,
        )
        self.section = Section.objects.create(section_name='Cabinet Events A', subject_code='CAB-EVENT-A')
        self.student = User.objects.create_user(
            username='cabinet-event-student',
            email='cabinet-event-student@example.com',
            password='secret1234',
            role=User.RoleChoices.STUDENT,
            first_name='Avery',
            last_name='Student',
            student_id='CAB-EVENT-001',
            nfc_uid='CARD-CAB-EVENT-001',
            section=self.section,
        )
        self.other_student = User.objects.create_user(
            username='cabinet-event-other',
            email='cabinet-event-other@example.com',
            password='secret1234',
            role=User.RoleChoices.STUDENT,
            student_id='CAB-EVENT-002',
        )
        opened_at = timezone.now() - timedelta(seconds=12)
        self.session = CabinetSession.objects.create(station='Station 1', status=CabinetSession.StatusChoices.CLOSED)
        self.session.opened_at = opened_at
        self.session.closed_at = opened_at + timedelta(seconds=12)
        self.session.save(update_fields=['opened_at', 'closed_at'])
        self.event = CabinetEvent.objects.create(
            user=self.student,
            event_type='cabinet_opened',
            details={
                'event_type': 'cabinet_opened',
                'station': 'Station 1',
                'cabinet_id': 'Cabinet A',
                'nfc_uid': self.student.nfc_uid,
                'access_method': 'NFC',
                'result': 'success',
                'cabinet_session_id': self.session.pk,
            },
        )
        self.other_event = CabinetEvent.objects.create(
            user=self.other_student,
            event_type='access_denied',
            details={'event_type': 'access_denied', 'result': 'rejected', 'failure_reason': 'Inactive account'},
        )

    def test_user_filter_and_event_details_use_real_owner_relationship(self):
        self.client.force_authenticate(user=self.admin)

        response = self.client.get(f'/api/cabinet-events/?user={self.student.pk}')

        self.assertEqual(response.status_code, 200, response.json())
        payload = response.json()['results']
        self.assertEqual([entry['id'] for entry in payload], [self.event.pk])
        self.assertEqual(payload[0]['user'], self.student.pk)
        self.assertEqual(payload[0]['user_name'], 'Avery Student')
        self.assertEqual(payload[0]['user_identifier'], 'CAB-EVENT-001')
        self.assertEqual(payload[0]['role'], User.RoleChoices.STUDENT)
        self.assertEqual(payload[0]['section_name'], 'Cabinet Events A')
        self.assertEqual(payload[0]['station'], 'Station 1')
        self.assertEqual(payload[0]['cabinet_id'], 'Cabinet A')
        self.assertEqual(payload[0]['event_type'], 'Cabinet Opened')
        self.assertEqual(payload[0]['result'], 'Success')
        self.assertEqual(payload[0]['access_method'], 'NFC')
        self.assertEqual(payload[0]['nfc_uid'], self.student.nfc_uid)
        self.assertEqual(payload[0]['duration_seconds'], 12)

        detail_response = self.client.get(f'/api/cabinet-events/{self.event.pk}/')
        self.assertEqual(detail_response.status_code, 200, detail_response.json())
        self.assertEqual(detail_response.json()['id'], self.event.pk)

        denied_response = self.client.get(f'/api/cabinet-events/?user={self.other_student.pk}')
        self.assertEqual(denied_response.status_code, 200, denied_response.json())
        denied_event = denied_response.json()['results'][0]
        self.assertEqual(denied_event['user'], self.other_student.pk)
        self.assertEqual(denied_event['event_type'], 'Access Denied')
        self.assertEqual(denied_event['failure_reason'], 'Inactive account')

    def test_event_records_cannot_be_edited_but_admin_can_delete(self):
        self.client.force_authenticate(user=self.admin)
        detail_url = f'/api/cabinet-events/{self.event.pk}/'

        update_response = self.client.patch(detail_url, {'event_type': 'access_denied'}, format='json')
        self.assertEqual(update_response.status_code, 405)

        self.client.force_authenticate(user=self.student)
        denied_delete = self.client.delete(detail_url)
        self.assertEqual(denied_delete.status_code, 403)

        self.client.force_authenticate(user=self.admin)
        delete_response = self.client.delete(detail_url)
        self.assertEqual(delete_response.status_code, 204)
        self.assertFalse(CabinetEvent.objects.filter(pk=self.event.pk).exists())

    def test_all_supported_cabinet_event_types_have_meaningful_results(self):
        expected = {
            'access_granted': ('Access Granted', 'Success'),
            'access_denied': ('Access Denied', 'Denied'),
            'cabinet_opened': ('Cabinet Opened', 'Success'),
            'cabinet_closed': ('Cabinet Closed', 'Success'),
            'unlock_failed': ('Unlock Failed', 'Failed'),
            'session_timeout': ('Session Timeout', 'Timeout'),
        }
        for event_type in expected:
            details = {'event_type': event_type}
            if event_type not in {'access_denied', 'unlock_failed', 'session_timeout'}:
                details['result'] = 'success'
            if event_type == 'unlock_failed':
                details['failure_reason'] = 'Cabinet unlock failed'
            CabinetEvent.objects.create(user=self.student, event_type=event_type, details=details)

        self.client.force_authenticate(user=self.admin)
        response = self.client.get(f'/api/cabinet-events/?user={self.student.pk}')
        self.assertEqual(response.status_code, 200, response.json())
        serialized = {entry['details']['event_type']: (entry['event_type'], entry['result']) for entry in response.json()['results']}
        for event_type, expected_values in expected.items():
            self.assertEqual(serialized[event_type], expected_values)


class UserProfileAccessLogTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.admin = User.objects.create_user(
            username='profile-log-admin',
            email='profile-log-admin@example.com',
            password='secret1234',
            role=User.RoleChoices.ADMIN,
        )
        self.instructor = User.objects.create_user(
            username='profile-log-instructor',
            email='profile-log-instructor@example.com',
            password='secret1234',
            role=User.RoleChoices.INSTRUCTOR,
            instructor_id='PROFILE-LOG-INSTRUCTOR',
        )
        self.section = Section.objects.create(
            section_name='Profile Log Section',
            subject_code='PROFILE-LOG',
            instructor=self.instructor,
        )
        self.mary_section = Section.objects.create(
            section_name='Profile Log Section B',
            subject_code='PROFILE-LOG-B',
        )
        self.john = User.objects.create_user(
            username='profile-log-john',
            email='john@example.com',
            password='secret1234',
            role=User.RoleChoices.STUDENT,
            student_id='PROFILE-LOG-JOHN',
            section=self.section,
        )
        self.mary = User.objects.create_user(
            username='profile-log-mary',
            email='mary@example.com',
            password='secret1234',
            role=User.RoleChoices.STUDENT,
            student_id='PROFILE-LOG-MARY',
            section=self.mary_section,
        )
        self.john_logs = [
            AccessLog.objects.create(user=self.john, action='open', status='success', station='Station 1'),
            AccessLog.objects.create(user=self.john, action='close', status='success', station='Station 1'),
            AccessLog.objects.create(user=self.john, action='open', status='success', station='Station 2'),
        ]
        self.mary_logs = [
            AccessLog.objects.create(user=self.mary, action='open', status='success', station='Station 2'),
            AccessLog.objects.create(user=self.mary, action='close', status='success', station='Station 2'),
        ]

    def test_admin_profile_logs_are_user_specific_and_global_logs_remain_global(self):
        self.client.force_authenticate(user=self.admin)

        for user, expected_logs in ((self.john, self.john_logs), (self.mary, self.mary_logs)):
            response = self.client.get(f'/api/users/{user.pk}/access-logs/?page_size=5')
            self.assertEqual(response.status_code, 200, response.json())
            self.assertEqual(
                {entry['id'] for entry in response.json()['results']},
                {log.pk for log in expected_logs},
            )
            self.assertEqual({entry['user'] for entry in response.json()['results']}, {user.pk})

        global_response = self.client.get('/api/access-logs/')
        self.assertEqual(global_response.status_code, 200, global_response.json())
        global_users = {entry['user'] for entry in global_response.json()['results']}
        self.assertEqual(global_users, {self.john.pk, self.mary.pk})

    def test_admin_profile_details_scope_records_and_counts_to_selected_user(self):
        john_activity = Activity.objects.create(title='John activity', created_by=self.instructor, section=self.section)
        mary_activity = Activity.objects.create(title='Mary activity', created_by=self.instructor, section=self.mary_section)
        john_submission = Submission.objects.create(activity=john_activity, student=self.john)
        mary_submission = Submission.objects.create(activity=mary_activity, student=self.mary)
        self.client.force_authenticate(user=self.admin)

        for user, own_logs, own_activity, own_submission in (
            (self.john, self.john_logs, john_activity, john_submission),
            (self.mary, self.mary_logs, mary_activity, mary_submission),
        ):
            response = self.client.get(f'/api/users/{user.pk}/profile-details/')
            self.assertEqual(response.status_code, 200, response.json())
            payload = response.json()
            self.assertEqual(payload['counts'], {
                'access_logs': len(own_logs),
                'activities': 1,
                'submissions': 1,
            })
            self.assertEqual(
                {entry['id'] for entry in payload['access_logs']},
                {log.pk for log in own_logs},
            )
            self.assertEqual({entry['id'] for entry in payload['activities']}, {own_activity.pk})
            self.assertEqual({entry['id'] for entry in payload['submissions']}, {own_submission.pk})
            for endpoint, expected_id in (
                ('activities', own_activity.pk),
                ('submissions', own_submission.pk),
            ):
                list_response = self.client.get(f'/api/users/{user.pk}/{endpoint}/')
                self.assertEqual(list_response.status_code, 200, list_response.json())
                self.assertEqual(
                    {entry['id'] for entry in list_response.json()['results']},
                    {expected_id},
                )

    def test_access_log_detail_returns_exact_cabinet_event_fields(self):
        opened_at = timezone.now() - timedelta(minutes=5)
        session = CabinetSession.objects.create(
            station='Station 4',
            status=CabinetSession.StatusChoices.CLOSED,
        )
        session.opened_at = opened_at
        session.closed_at = opened_at + timedelta(minutes=5)
        session.save(update_fields=['opened_at', 'closed_at'])
        log = AccessLog.objects.create(
            user=self.john,
            cabinet_session=session,
            action=AccessLog.AccessActionChoices.OPEN,
            status=AccessLog.AccessStatusChoices.SUCCESS,
            nfc_uid='ABCD1234',
            reason='Cabinet session opened',
        )
        self.client.force_authenticate(user=self.admin)

        response = self.client.get(f'/api/access-logs/{log.pk}/')

        self.assertEqual(response.status_code, 200, response.json())
        payload = response.json()
        self.assertEqual(payload['id'], log.pk)
        self.assertEqual(payload['user'], self.john.pk)
        self.assertEqual(payload['username'], self.john.username)
        self.assertEqual(payload['role'], User.RoleChoices.STUDENT)
        self.assertEqual(payload['station'], 'Station 4')
        self.assertEqual(payload['access_type'], 'Cabinet Opened')
        self.assertEqual(payload['access_method'], 'NFC')
        self.assertEqual(payload['nfc_uid'], 'ABCD1234')
        self.assertEqual(payload['status'], AccessLog.AccessStatusChoices.SUCCESS)
        self.assertEqual(payload['duration_seconds'], 300)

    def test_instructor_and_student_profile_log_requests_obey_user_scope(self):
        self.client.force_authenticate(user=self.instructor)
        instructor_response = self.client.get(f'/api/users/{self.john.pk}/access-logs/')
        self.assertEqual(instructor_response.status_code, 200, instructor_response.json())
        self.assertEqual(
            {entry['user'] for entry in instructor_response.json()['results']},
            {self.john.pk},
        )
        self.assertEqual(self.client.get(f'/api/users/{self.mary.pk}/access-logs/').status_code, 404)
        self.assertEqual(self.client.get(f'/api/access-logs/{self.mary_logs[0].pk}/').status_code, 404)
        self.assertEqual(self.client.get(f'/api/users/{self.john.pk}/profile-details/').status_code, 403)
        self.assertEqual(self.client.get(f'/api/users/{self.john.pk}/activities/').status_code, 403)
        self.assertEqual(self.client.get(f'/api/users/{self.john.pk}/submissions/').status_code, 403)

        self.client.force_authenticate(user=self.john)
        self_response = self.client.get(f'/api/users/{self.john.pk}/access-logs/')
        other_response = self.client.get(f'/api/users/{self.mary.pk}/access-logs/')
        self.assertEqual(self_response.status_code, 200, self_response.json())
        self.assertEqual({entry['user'] for entry in self_response.json()['results']}, {self.john.pk})
        self.assertEqual(other_response.status_code, 404)
        self.assertEqual(self.client.get(f'/api/access-logs/{self.john_logs[0].pk}/').status_code, 200)
        self.assertEqual(self.client.get(f'/api/access-logs/{self.mary_logs[0].pk}/').status_code, 404)
        self.assertEqual(self.client.get(f'/api/users/{self.john.pk}/profile-details/').status_code, 403)
        self.assertEqual(self.client.get(f'/api/users/{self.john.pk}/activities/').status_code, 403)
        self.assertEqual(self.client.get(f'/api/users/{self.john.pk}/submissions/').status_code, 403)


@override_settings(TAPTRACK_NFC_DEVICE_MAP=[
    {'device_id': 'PHOTO-CABINET', 'api_key': 'photo-test-key', 'station': 'Station 1', 'cabinet_name': 'Cabinet 1', 'active': True},
])
class AccessLogPhotoApiTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.device_api_key = 'photo-test-key'
        os.environ['DEVICE_API_KEY'] = self.device_api_key
        self.media_directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.media_directory.cleanup)
        self.enterContext(override_settings(MEDIA_ROOT=self.media_directory.name))
        self.student = User.objects.create_user(
            username='photo-student',
            email='photo-student@example.com',
            password='secret1234',
            role=User.RoleChoices.STUDENT,
            student_id='PHOTO-001',
        )
        self.access_log = AccessLog.objects.create(
            user=self.student,
            nfc_uid='PHOTO-UID-001',
            station='Station 1',
            cabinet_name='Cabinet 1',
            status=AccessLog.AccessStatusChoices.SUCCESS,
        )

    def jpeg_upload(self, filename='camera-supplied-name.jpg'):
        image_data = BytesIO()
        Image.new('RGB', (8, 8), color=(30, 80, 120)).save(image_data, format='JPEG')
        return SimpleUploadedFile(filename, image_data.getvalue(), content_type='image/jpeg')

    def upload_url(self, event_id=None):
        return f'/api/access-logs/{event_id or self.access_log.pk}/photo/'

    def test_authorized_device_uploads_jpeg_to_existing_event(self):
        response = self.client.post(
            self.upload_url(),
            {'image': self.jpeg_upload(), 'station': 'Forged station', 'student_id': 'FORGED'},
            HTTP_X_API_KEY=self.device_api_key,
            format='multipart',
        )

        self.assertEqual(response.status_code, 201, response.json())
        photo = AccessLogPhoto.objects.get(access_log=self.access_log)
        self.assertEqual(photo.capture_status, AccessLogPhoto.CaptureStatusChoices.SUCCESS)
        self.assertIsNotNone(photo.captured_at)
        self.assertTrue(photo.image.name.startswith('access-log-photos/'))
        self.assertTrue(photo.image.name.endswith('.jpg'))
        self.assertNotIn('camera-supplied-name', photo.image.name)
        self.assertEqual(response.json()['event_id'], self.access_log.pk)
        self.assertNotIn('url', response.json())
        self.assertEqual(photo.access_log.station, 'Station 1')
        self.assertEqual(photo.access_log.user, self.student)

    def test_photo_metadata_reports_unavailable_for_access_log_without_photo(self):
        response = self.client.get(
            self.upload_url(), HTTP_X_API_KEY=self.device_api_key,
        )

        self.assertEqual(response.status_code, 200, response.json())
        self.assertEqual(response.json(), {
            'event_id': self.access_log.pk,
            'capture_status': None,
            'captured_at': None,
            'diagnostic_error': None,
            'image_available': False,
            'image_endpoint': None,
        })
        missing_image = self.client.get(
            f'/api/access-logs/{self.access_log.pk}/photo/image/',
            HTTP_X_API_KEY=self.device_api_key,
        )
        self.assertEqual(missing_image.status_code, 404)

    def test_pending_and_failed_capture_statuses_are_read_back_from_backend(self):
        pending = self.client.post(
            f'/api/access-logs/{self.access_log.pk}/photo-status/',
            {'status': 'pending'}, HTTP_X_API_KEY=self.device_api_key, format='json',
        )
        pending_metadata = self.client.get(self.upload_url(), HTTP_X_API_KEY=self.device_api_key)
        self.assertEqual(pending.status_code, 200, pending.json())
        self.assertEqual(pending_metadata.json()['capture_status'], 'pending')
        self.assertIsNone(pending_metadata.json()['captured_at'])
        self.assertFalse(pending_metadata.json()['image_available'])

        failed = self.client.post(
            f'/api/access-logs/{self.access_log.pk}/photo-status/',
            {'status': 'failed', 'error': 'Camera unavailable'},
            HTTP_X_API_KEY=self.device_api_key, format='json',
        )
        failed_metadata = self.client.get(self.upload_url(), HTTP_X_API_KEY=self.device_api_key)
        self.assertEqual(failed.status_code, 200, failed.json())
        self.assertEqual(failed_metadata.json()['capture_status'], 'failed')
        self.assertEqual(failed_metadata.json()['diagnostic_error'], 'Camera unavailable')
        self.assertIsNotNone(failed_metadata.json()['captured_at'])

    def test_same_cabinet_device_can_read_a_station_two_event(self):
        station_two_event = AccessLog.objects.create(
            user=self.student,
            nfc_uid='PHOTO-STATION-TWO',
            station='Station 2',
            cabinet_name='Cabinet 1',
            status=AccessLog.AccessStatusChoices.SUCCESS,
        )
        uploaded = self.client.post(
            self.upload_url(station_two_event.pk), {'image': self.jpeg_upload()},
            HTTP_X_API_KEY=self.device_api_key, format='multipart',
        )
        metadata = self.client.get(
            self.upload_url(station_two_event.pk), HTTP_X_API_KEY=self.device_api_key,
        )

        self.assertEqual(uploaded.status_code, 201, uploaded.json())
        self.assertEqual(metadata.status_code, 200, metadata.json())
        self.assertEqual(metadata.json()['event_id'], station_two_event.pk)
        self.assertEqual(metadata.json()['capture_status'], 'success')
        image = self.client.get(metadata.json()['image_endpoint'], HTTP_X_API_KEY=self.device_api_key)
        self.assertEqual(image.status_code, 200)
        image.close()

    def test_missing_and_invalid_device_credentials_cannot_read_photo_data(self):
        for path in (self.upload_url(), f'/api/access-logs/{self.access_log.pk}/photo/image/'):
            with self.subTest(path=path):
                missing = self.client.get(path)
                invalid = self.client.get(path, HTTP_X_API_KEY='invalid-photo-key')
                self.assertEqual(missing.status_code, 401)
                self.assertEqual(invalid.status_code, 401)

    def test_photo_metadata_and_image_bytes_require_authorized_device(self):
        uploaded = self.client.post(
            self.upload_url(), {'image': self.jpeg_upload()},
            HTTP_X_API_KEY=self.device_api_key, format='multipart',
        )
        self.assertEqual(uploaded.status_code, 201, uploaded.json())

        metadata = self.client.get(self.upload_url(), HTTP_X_API_KEY=self.device_api_key)
        self.assertEqual(metadata.status_code, 200, metadata.json())
        self.assertEqual(metadata.json()['capture_status'], AccessLogPhoto.CaptureStatusChoices.SUCCESS)
        self.assertTrue(metadata.json()['image_available'])
        self.assertEqual(
            metadata.json()['image_endpoint'],
            f'/api/access-logs/{self.access_log.pk}/photo/image/',
        )
        self.assertNotIn('/media/', metadata.json()['image_endpoint'])

        image = self.client.get(metadata.json()['image_endpoint'], HTTP_X_API_KEY=self.device_api_key)
        self.assertEqual(image.status_code, 200)
        self.assertEqual(image['Content-Type'], 'image/jpeg')
        self.assertIn('no-store', image['Cache-Control'])
        image_bytes = b''.join(image.streaming_content)
        self.assertTrue(image_bytes.startswith(b'\xff\xd8'))
        image.close()

        missing_key = self.client.get(self.upload_url())
        self.assertEqual(missing_key.status_code, 401)

    def test_device_photo_reads_are_limited_to_its_cabinet(self):
        other_cabinet_event = AccessLog.objects.create(
            nfc_uid='PHOTO-OTHER-CABINET', station='Station 2', cabinet_name='Cabinet 2',
        )

        response = self.client.get(
            self.upload_url(other_cabinet_event.pk), HTTP_X_API_KEY=self.device_api_key,
        )

        self.assertEqual(response.status_code, 404)

    def test_jwt_photo_reads_follow_student_and_instructor_log_permissions(self):
        uploaded = self.client.post(
            self.upload_url(), {'image': self.jpeg_upload()},
            HTTP_X_API_KEY=self.device_api_key, format='multipart',
        )
        self.assertEqual(uploaded.status_code, 201, uploaded.json())
        image_path = f'/api/access-logs/{self.access_log.pk}/photo/image/'
        other_student = User.objects.create_user(
            username='photo-other-student', email='photo-other@example.com',
            password='secret1234', role=User.RoleChoices.STUDENT, student_id='PHOTO-002',
        )
        another_cabinet_log = AccessLog.objects.create(
            user=other_student,
            nfc_uid='PHOTO-ADMIN-OTHER-CABINET',
            station='Station 2',
            cabinet_name='Cabinet 2',
            status=AccessLog.AccessStatusChoices.SUCCESS,
        )
        another_cabinet_photo_path = f'/api/access-logs/{another_cabinet_log.pk}/photo/'
        self.client.force_authenticate(user=other_student)
        self.assertEqual(self.client.get(another_cabinet_photo_path).status_code, 200)
        self.assertEqual(self.client.get(f'{another_cabinet_photo_path}image/').status_code, 403)

        self.client.force_authenticate(user=self.student)
        self.assertEqual(self.client.get(self.upload_url()).status_code, 200)
        student_logs = self.client.get('/api/access-logs/').json()['results']
        self.assertEqual({entry['user'] for entry in student_logs}, {self.student.pk})
        self.assertEqual(student_logs[0]['photo_capture_status'], 'success')
        self.assertNotIn('image_endpoint', student_logs[0])
        self.assertEqual(self.client.get(another_cabinet_photo_path).status_code, 404)
        self.assertEqual(self.client.get(f'{another_cabinet_photo_path}image/').status_code, 404)
        student_metadata = self.client.get(self.upload_url()).json()
        self.assertEqual(student_metadata['capture_status'], 'success')
        self.assertIsNone(student_metadata['captured_at'])
        self.assertIsNone(student_metadata['diagnostic_error'])
        self.assertFalse(student_metadata['image_available'])
        self.assertIsNone(student_metadata['image_endpoint'])
        student_image = self.client.get(image_path)
        self.assertEqual(student_image.status_code, 403)

        instructor = User.objects.create_user(
            username='photo-instructor', email='photo-instructor@example.com',
            password='secret1234', role=User.RoleChoices.INSTRUCTOR, instructor_id='PHOTO-INS-1',
        )
        section = Section.objects.create(
            section_name='Photo Access Section', subject_code='PHOTO-101', instructor=instructor,
        )
        self.student.section = section
        self.student.save(update_fields=['section'])
        self.client.force_authenticate(user=instructor)

        self.assertEqual(self.client.get(self.upload_url()).status_code, 200)
        instructor_image = self.client.get(image_path)
        self.assertEqual(instructor_image.status_code, 200)
        instructor_image.close()

        assigned_section = Section.objects.create(
            section_name='M2M Assigned Photo Section', subject_code='PHOTO-ASSIGNED',
        )
        assigned_section.assigned_instructors.add(instructor)
        assigned_student = User.objects.create_user(
            username='photo-m2m-assigned-student', email='photo-m2m-assigned@example.com',
            password='secret1234', role=User.RoleChoices.STUDENT, student_id='PHOTO-ASSIGNED-1',
            section=assigned_section,
        )
        assigned_log = AccessLog.objects.create(
            user=assigned_student, nfc_uid='PHOTO-M2M-ASSIGNED', station='Station 2',
            cabinet_name='Cabinet 1', status=AccessLog.AccessStatusChoices.SUCCESS,
        )
        self.client.force_authenticate(user=instructor)
        self.assertEqual(self.client.get(f'/api/access-logs/{assigned_log.pk}/').status_code, 200)
        self.assertEqual(self.client.get(f'/api/access-logs/{assigned_log.pk}/photo/').status_code, 200)

        assigned_upload = self.client.post(
            self.upload_url(assigned_log.pk), {'image': self.jpeg_upload()},
            HTTP_X_API_KEY=self.device_api_key, format='multipart',
        )
        self.assertEqual(assigned_upload.status_code, 201, assigned_upload.json())
        assigned_image = self.client.get(f'/api/access-logs/{assigned_log.pk}/photo/image/')
        self.assertEqual(assigned_image.status_code, 200)
        assigned_image.close()

        unrelated_instructor = User.objects.create_user(
            username='unrelated-photo-instructor', email='unrelated-photo@example.com',
            password='secret1234', role=User.RoleChoices.INSTRUCTOR, instructor_id='PHOTO-INS-2',
        )
        Section.objects.create(
            section_name='Unrelated Photo Section', subject_code='PHOTO-202', instructor=unrelated_instructor,
        )
        self.client.force_authenticate(user=unrelated_instructor)
        self.assertEqual(self.client.get(self.upload_url()).status_code, 404)
        self.assertEqual(self.client.get(image_path).status_code, 404)

        admin_user = User.objects.create_superuser(
            username='photo-read-admin', email='photo-read-admin@example.com', password='secret1234',
        )
        self.client.force_authenticate(user=admin_user)
        self.assertEqual(self.client.get(self.upload_url()).status_code, 200)
        self.assertEqual(self.client.get(another_cabinet_photo_path).status_code, 200)
        admin_image = self.client.get(image_path)
        self.assertEqual(admin_image.status_code, 200)
        admin_image.close()

    def test_upload_requires_valid_device_api_key(self):
        missing = self.client.post(self.upload_url(), {'image': self.jpeg_upload()}, format='multipart')
        invalid = self.client.post(
            self.upload_url(), {'image': self.jpeg_upload()},
            HTTP_X_API_KEY='incorrect-key', format='multipart',
        )

        self.assertEqual(missing.status_code, 403)
        self.assertEqual(invalid.status_code, 403)
        self.assertFalse(AccessLogPhoto.objects.exists())

    def test_upload_rejects_invalid_image_and_oversized_file(self):
        invalid = SimpleUploadedFile('not-an-image.jpg', b'not a jpeg', content_type='image/jpeg')
        invalid_response = self.client.post(
            self.upload_url(), {'image': invalid}, HTTP_X_API_KEY=self.device_api_key, format='multipart',
        )
        oversized = SimpleUploadedFile(
            'large.jpg', b'x' * (5 * 1024 * 1024 + 1), content_type='image/jpeg',
        )
        oversized_response = self.client.post(
            self.upload_url(), {'image': oversized}, HTTP_X_API_KEY=self.device_api_key, format='multipart',
        )

        self.assertEqual(invalid_response.status_code, 400)
        self.assertEqual(oversized_response.status_code, 413)
        self.assertFalse(AccessLogPhoto.objects.exists())

    def test_device_cannot_upload_to_another_cabinets_event(self):
        other_cabinet_event = AccessLog.objects.create(
            nfc_uid='OTHER-CABINET-UID', station='Station 2', cabinet_name='Cabinet 2',
        )
        response = self.client.post(
            self.upload_url(other_cabinet_event.pk), {'image': self.jpeg_upload()},
            HTTP_X_API_KEY=self.device_api_key, format='multipart',
        )

        self.assertEqual(response.status_code, 404)
        self.assertFalse(AccessLogPhoto.objects.exists())

    def test_duplicate_photo_upload_is_rejected(self):
        first = self.client.post(
            self.upload_url(), {'image': self.jpeg_upload()},
            HTTP_X_API_KEY=self.device_api_key, format='multipart',
        )
        duplicate = self.client.post(
            self.upload_url(), {'image': self.jpeg_upload()},
            HTTP_X_API_KEY=self.device_api_key, format='multipart',
        )

        self.assertEqual(first.status_code, 201, first.json())
        self.assertEqual(duplicate.status_code, 409, duplicate.json())
        self.assertEqual(AccessLogPhoto.objects.filter(access_log=self.access_log).count(), 1)

    def test_camera_failure_status_is_recorded_without_changing_access_event(self):
        response = self.client.post(
            f'/api/access-logs/{self.access_log.pk}/photo-status/',
            {'status': 'failed', 'error': 'Camera did not respond'},
            HTTP_X_API_KEY=self.device_api_key,
            format='json',
        )

        self.assertEqual(response.status_code, 200, response.json())
        photo = AccessLogPhoto.objects.get(access_log=self.access_log)
        self.assertEqual(photo.capture_status, AccessLogPhoto.CaptureStatusChoices.FAILED)
        self.assertEqual(photo.diagnostic_error, 'Camera did not respond')
        self.assertIsNotNone(photo.captured_at)
        self.access_log.refresh_from_db()
        self.assertEqual(self.access_log.status, AccessLog.AccessStatusChoices.SUCCESS)

    def test_admin_displays_photo_and_private_file_is_staff_only(self):
        uploaded = self.client.post(
            self.upload_url(), {'image': self.jpeg_upload()},
            HTTP_X_API_KEY=self.device_api_key, format='multipart',
        )
        self.assertEqual(uploaded.status_code, 201, uploaded.json())
        admin_user = User.objects.create_superuser(
            username='photo-admin', email='photo-admin@example.com', password='secret1234',
        )
        self.client.force_login(admin_user)

        listing = self.client.get(reverse('admin:api_accesslog_changelist'))
        self.assertEqual(listing.status_code, 200)
        self.assertContains(listing, '<img', html=False)
        photo = AccessLogPhoto.objects.get(access_log=self.access_log)
        protected = self.client.get(reverse('admin:api_accesslog_photo', args=[self.access_log.pk]))
        public_media = self.client.get(f'/media/{photo.image.name}')
        self.assertEqual(protected.status_code, 200)
        self.assertEqual(protected['Content-Type'], 'image/jpeg')
        self.assertIn('no-store', protected['Cache-Control'])
        self.assertEqual(public_media.status_code, 404)
        protected.close()
        self.client.logout()
        unauthenticated = self.client.get(reverse('admin:api_accesslog_photo', args=[self.access_log.pk]))
        self.assertEqual(unauthenticated.status_code, 302)

    def test_legacy_access_log_without_photo_remains_valid(self):
        access_log_admin = AccessLogAdmin(model=AccessLog, admin_site=django_admin.site)

        self.assertFalse(AccessLogPhoto.objects.filter(access_log=self.access_log).exists())
        self.assertEqual(access_log_admin.photo_thumbnail(self.access_log), 'No photo')
        self.assertEqual(access_log_admin.photo_capture_status(self.access_log), 'Not captured')
        self.assertEqual(self.access_log.status, AccessLog.AccessStatusChoices.SUCCESS)

    def test_deleting_photo_metadata_removes_stored_image(self):
        response = self.client.post(
            self.upload_url(), {'image': self.jpeg_upload()},
            HTTP_X_API_KEY=self.device_api_key, format='multipart',
        )
        self.assertEqual(response.status_code, 201, response.json())
        photo = AccessLogPhoto.objects.get(access_log=self.access_log)
        stored_name = photo.image.name
        self.assertTrue(photo.image.storage.exists(stored_name))

        photo.delete()

        self.assertFalse(photo.image.storage.exists(stored_name))


class AccessLogPhotoMigrationTests(TransactionTestCase):
    migrate_from = ('api', '0035_cabinetsession_unique_open_station')
    migrate_to = ('api', '0036_accesslogphoto')

    def setUp(self):
        executor = MigrationExecutor(connection)
        executor.migrate([self.migrate_from])
        old_apps = executor.loader.project_state([self.migrate_from]).apps
        historical_log = old_apps.get_model('api', 'AccessLog').objects.create(
            nfc_uid='MIGRATION-LEGACY-UID',
            station='Station 1',
            cabinet_name='Cabinet 1',
            status='success',
        )
        self.access_log_id = historical_log.pk
        MigrationExecutor(connection).migrate([self.migrate_to])

    def tearDown(self):
        executor = MigrationExecutor(connection)
        executor.migrate(executor.loader.graph.leaf_nodes())

    def test_existing_access_log_survives_photo_migration_without_photo(self):
        new_apps = MigrationExecutor(connection).loader.project_state([self.migrate_to]).apps
        historical_log = new_apps.get_model('api', 'AccessLog').objects.get(pk=self.access_log_id)
        photo_model = new_apps.get_model('api', 'AccessLogPhoto')

        self.assertEqual(historical_log.nfc_uid, 'MIGRATION-LEGACY-UID')
        self.assertEqual(historical_log.station, 'Station 1')
        self.assertEqual(historical_log.status, 'success')
        self.assertFalse(photo_model.objects.filter(access_log_id=self.access_log_id).exists())


@override_settings(TAPTRACK_NFC_DEVICE_MAP=[
    {'device_id': 'CABINET1-STATION1', 'api_key': 'cabinet-device-test-key', 'station': 'Station 1', 'cabinet_name': 'Cabinet 1', 'active': True},
])
class CabinetDeviceIntegrationTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.device_api_key = 'cabinet-device-test-key'
        os.environ['DEVICE_API_KEY'] = self.device_api_key
        self.section = Section.objects.create(section_name='Cabinet Section', subject_code='CAB-101')

    @override_settings(TAPTRACK_CABINET_STATIONS=['Station 1', 'Station 2'])
    def test_station_status_is_derived_per_station_and_hides_details_from_students(self):
        student = User.objects.create_user(
            username='station-status-student',
            email='station-status@example.com',
            password='secret1234',
            role=User.RoleChoices.STUDENT,
            student_id='STATION-STATUS-001',
        )
        session = CabinetSession.objects.create(
            station='Station 1',
            status=CabinetSession.StatusChoices.OPEN,
        )
        session.opened_by.add(student)
        self.client.force_authenticate(user=student)

        response = self.client.get('/api/cabinet-sessions/station-status/')

        self.assertEqual(response.status_code, 200, response.json())
        stations = response.json()['stations']
        self.assertEqual(
            [(item['station'], item['status']) for item in stations],
            [('Station 1', 'OCCUPIED'), ('Station 2', 'AVAILABLE')],
        )
        self.assertIsNone(stations[0]['active_session'])
        station_two_session = CabinetSession.objects.create(
            station='Station 2',
            status=CabinetSession.StatusChoices.OPEN,
        )

        instructor = User.objects.create_user(
            username='station-status-instructor',
            email='station-status-instructor@example.com',
            password='secret1234',
            role=User.RoleChoices.INSTRUCTOR,
            instructor_id='STATION-STATUS-INSTRUCTOR',
        )
        self.client.force_authenticate(user=instructor)
        instructor_response = self.client.get('/api/cabinet-sessions/station-status/')
        instructor_station = instructor_response.json()['stations'][0]
        self.assertEqual(instructor_station['status'], 'OCCUPIED')
        self.assertEqual(instructor_station['active_session']['participant_count'], 1)
        self.assertTrue(instructor_station['active_session']['opened_at'])
        self.assertEqual(instructor_response.json()['stations'][1]['status'], 'OCCUPIED')

        admin = User.objects.create_user(
            username='station-status-admin',
            email='station-status-admin@example.com',
            password='secret1234',
            role=User.RoleChoices.ADMIN,
        )
        self.client.force_authenticate(user=admin)
        admin_response = self.client.get('/api/cabinet-sessions/station-status/')
        admin_station = admin_response.json()['stations'][0]
        self.assertEqual(admin_station['active_session']['id'], session.pk)
        self.assertEqual(admin_station['active_session']['station'], 'Station 1')

        session.status = CabinetSession.StatusChoices.CLOSED
        session.save(update_fields=['status', 'updated_at'])
        closed_response = self.client.get('/api/cabinet-sessions/station-status/')
        self.assertEqual(closed_response.json()['stations'][0]['status'], 'AVAILABLE')
        self.assertEqual(closed_response.json()['stations'][1]['status'], 'OCCUPIED')

        station_two_session.status = CabinetSession.StatusChoices.CLOSED
        station_two_session.save(update_fields=['status', 'updated_at'])
        all_available_response = self.client.get('/api/cabinet-sessions/station-status/')
        self.assertEqual(
            [item['status'] for item in all_available_response.json()['stations']],
            ['AVAILABLE', 'AVAILABLE'],
        )

    @override_settings(
        DEBUG=True,
        TAPTRACK_NFC_DEVICE_MAP=[{
            'device_id': 'CABINET1-STATION1',
            'api_key': 'cabinet-device-test-key',
            'station': 'Station 1',
            'cabinet_name': 'Cabinet 1',
            'active': True,
        }],
        TAPTRACK_CABINET_STATIONS=['Station 1', 'Station 2'],
    )
    def test_open_workflow_rejects_same_station_but_allows_another_station(self):
        first = self.client.post(
            '/api/cabinet/workflow/',
            {'command': 'open', 'station': 'Station 1'},
            HTTP_X_API_KEY=self.device_api_key,
            HTTP_X_TAPTRACK_MOCK_MODE='true',
            format='json',
        )
        second_station = self.client.post(
            '/api/cabinet/workflow/',
            {'command': 'open', 'station': 'Station 2'},
            HTTP_X_API_KEY=self.device_api_key,
            HTTP_X_TAPTRACK_MOCK_MODE='true',
            format='json',
        )
        duplicate = self.client.post(
            '/api/cabinet/workflow/',
            {'command': 'open', 'station': 'Station 1'},
            HTTP_X_API_KEY=self.device_api_key,
            HTTP_X_TAPTRACK_MOCK_MODE='true',
            format='json',
        )

        self.assertEqual(first.status_code, 200, first.json())
        self.assertEqual(second_station.status_code, 200, second_station.json())
        self.assertEqual(duplicate.status_code, 409, duplicate.json())

    @override_settings(
        DEBUG=True,
        TAPTRACK_NFC_DEVICE_MAP=[{
            'device_id': 'CABINET1-STATION1',
            'api_key': 'cabinet-device-test-key',
            'station': 'Station 1',
            'cabinet_name': 'Cabinet 1',
            'active': True,
        }],
        TAPTRACK_CABINET_NAME='Cabinet 1',
        TAPTRACK_CABINET_STATIONS=['Station 1', 'Station 2'],
    )
    def test_mock_workflow_lists_configured_station_state_and_accepts_only_configured_choice(self):
        CabinetSession.objects.create(station='Station 2', status=CabinetSession.StatusChoices.OPEN)

        response = self.client.get(
            '/api/cabinet/workflow/',
            HTTP_X_API_KEY=self.device_api_key,
            HTTP_X_TAPTRACK_MOCK_MODE='true',
        )

        self.assertEqual(response.status_code, 200, response.json())
        payload = response.json()
        self.assertEqual(payload['mode'], 'mock')
        self.assertEqual([station['name'] for station in payload['stations']], ['Station 1', 'Station 2'])
        self.assertEqual([station['status'] for station in payload['stations']], ['available', 'occupied'])
        self.assertNotIn(self.device_api_key, response.content.decode())

        opened = self.client.post(
            '/api/cabinet/workflow/',
            {'command': 'open', 'station': 'Station 1'},
            HTTP_X_API_KEY=self.device_api_key,
            HTTP_X_TAPTRACK_MOCK_MODE='true',
            format='json',
        )
        self.assertEqual(opened.status_code, 200, opened.json())
        self.assertEqual(opened.json()['station'], 'Station 1')

        student = User.objects.create_user(
            username='mock-station-student',
            email='mock-station@example.com',
            password='secret1234',
            role=User.RoleChoices.STUDENT,
            student_id='MOCK-STATION-001',
            nfc_uid='MOCK-STATION-UID',
        )
        scan = self.client.post(
            '/api/verify-nfc/',
            {'nfc_uid': ' mock-station-uid ', 'station': 'Station 1'},
            HTTP_X_API_KEY=self.device_api_key,
            HTTP_X_TAPTRACK_MOCK_MODE='true',
            format='json',
        )
        self.assertEqual(scan.status_code, 200, scan.json())
        access_log = student.access_logs.get()
        self.assertEqual(access_log.cabinet_session_id, opened.json()['session']['id'])
        self.assertEqual(access_log.action, AccessLog.AccessActionChoices.OPEN)
        self.assertEqual(access_log.nfc_uid, 'MOCK-STATION-UID')
        self.assertEqual(access_log.station, 'Station 1')
        self.assertEqual(access_log.cabinet_name, 'Cabinet 1')
        self.assertEqual(access_log.reason, 'NFC verified successfully')
        self.assertEqual(scan.json()['event_id'], access_log.id)
        self.assertEqual(scan.json()['user']['student_id'], student.student_id)
        self.assertEqual(scan.json()['station']['name'], 'Station 1')

        second_student = User.objects.create_user(
            username='mock-station-student-two',
            email='mock-station-two@example.com',
            password='secret1234',
            role=User.RoleChoices.STUDENT,
            student_id='MOCK-STATION-002',
            nfc_uid='MOCK-STATION-UID-2',
        )
        second_scan = self.client.post(
            '/api/verify-nfc/',
            {'nfc_uid': second_student.nfc_uid, 'station': 'Station 1'},
            HTTP_X_API_KEY=self.device_api_key,
            HTTP_X_TAPTRACK_MOCK_MODE='true',
            format='json',
        )
        self.assertEqual(second_scan.status_code, 200, second_scan.json())

        opened_sessions = self.client.get(
            '/api/cabinet/workflow/',
            HTTP_X_API_KEY=self.device_api_key,
            HTTP_X_TAPTRACK_MOCK_MODE='true',
        ).json()['stations']
        opened_session = next(item['session'] for item in opened_sessions if item['name'] == 'Station 1')
        self.assertEqual({item['studentId'] for item in opened_session['opened_by']}, {student.student_id, second_student.student_id})
        self.assertEqual(opened_session['participant_count'], 2)
        self.assertNotIn('participants', opened_session)
        self.assertNotIn('closing_participants', opened_session)

        finish_open = self.client.post(
            '/api/cabinet/workflow/',
            {'command': 'finish_open', 'station': 'Station 1'},
            HTTP_X_API_KEY=self.device_api_key,
            HTTP_X_TAPTRACK_MOCK_MODE='true',
            format='json',
        )
        self.assertEqual(finish_open.status_code, 200, finish_open.json())
        start_close = self.client.post(
            '/api/cabinet/workflow/',
            {'command': 'start_close', 'station': 'Station 1'},
            HTTP_X_API_KEY=self.device_api_key,
            HTTP_X_TAPTRACK_MOCK_MODE='true',
            format='json',
        )
        self.assertEqual(start_close.status_code, 200, start_close.json())
        for closing_student in (student, second_student):
            closing_scan = self.client.post(
                '/api/verify-nfc/',
                {'nfc_uid': closing_student.nfc_uid, 'station': 'Station 1'},
                HTTP_X_API_KEY=self.device_api_key,
                HTTP_X_TAPTRACK_MOCK_MODE='true',
                format='json',
            )
            self.assertEqual(closing_scan.status_code, 200, closing_scan.json())

        closed_sessions = self.client.get(
            '/api/cabinet/workflow/',
            HTTP_X_API_KEY=self.device_api_key,
            HTTP_X_TAPTRACK_MOCK_MODE='true',
        ).json()['stations']
        closed_by_session = next(item['session'] for item in closed_sessions if item['name'] == 'Station 1')
        self.assertEqual({item['studentId'] for item in closed_by_session['closed_by']}, {student.student_id, second_student.student_id})
        self.assertEqual(closed_by_session['participant_count'], 2)

        invalid = self.client.post(
            '/api/cabinet/workflow/',
            {'command': 'open', 'station': 'Station 3'},
            HTTP_X_API_KEY=self.device_api_key,
            HTTP_X_TAPTRACK_MOCK_MODE='true',
            format='json',
        )
        self.assertEqual(invalid.status_code, 400)

    @override_settings(TAPTRACK_CABINET_STATIONS=[])
    def test_hardware_workflow_requires_a_configured_selected_station(self):
        student = User.objects.create_user(
            username='hardware-no-mock-station-student',
            email='hardware-no-mock-station@example.com',
            password='secret1234',
            role=User.RoleChoices.STUDENT,
            student_id='HARDWARE-NO-MOCK-001',
            nfc_uid='HARDWARE-READER-UID-001',
        )

        opened = self.client.post(
            '/api/cabinet/workflow/',
            {'command': 'open', 'station': 'Station 1'},
            HTTP_X_API_KEY=self.device_api_key,
            format='json',
        )
        self.assertEqual(opened.status_code, 400, opened.json())
        self.assertEqual(opened.json()['error'], 'Select a configured cabinet station.')
        self.assertFalse(student.access_logs.exists())

    @override_settings(DEBUG=True)
    def test_unmapped_hardware_key_is_not_silently_routed_to_mock_station_validation(self):
        os.environ['DEVICE_API_KEY'] = 'unmapped-hardware-key'

        workflow = self.client.get(
            '/api/cabinet/workflow/',
            HTTP_X_API_KEY='unmapped-hardware-key',
        )
        scan = self.client.post(
            '/api/verify-nfc/',
            {'nfc_uid': 'REAL-READER-UID'},
            HTTP_X_API_KEY='unmapped-hardware-key',
            format='json',
        )

        self.assertEqual(workflow.status_code, 403, workflow.json())
        self.assertEqual(scan.status_code, 403, scan.json())
        self.assertNotIn('mock station', workflow.json().get('error', '').lower())
        self.assertNotIn('mock station', scan.json().get('error', '').lower())
        self.assertFalse(AccessLog.objects.filter(nfc_uid='REAL-READER-UID').exists())

    def test_hardware_scan_uses_active_workflow_station_and_disabled_devices_are_rejected(self):
        student = User.objects.create_user(
            username='reader-bound-student',
            email='reader-bound@example.com',
            password='secret1234',
            role=User.RoleChoices.STUDENT,
            student_id='READER-BOUND-001',
            nfc_uid='READER-BOUND-UID',
        )

        station_one = CabinetSession.objects.create(
            station='Station 1',
            status=CabinetSession.StatusChoices.OPEN,
        )
        session = CabinetSession.objects.create(
            station='Station 2',
            status=CabinetSession.StatusChoices.OPEN,
            workflow_state=CabinetSession.WorkflowStateChoices.OPENING,
        )
        response = self.client.post(
            '/api/verify-nfc/',
            {'nfc_uid': student.nfc_uid, 'station': 'Station 2'},
            HTTP_X_API_KEY=self.device_api_key,
            format='json',
        )
        self.assertEqual(response.status_code, 200, response.json())
        log = AccessLog.objects.get(user=student)
        self.assertEqual(log.station, 'Station 2')
        self.assertEqual(log.cabinet_session, session)
        self.assertEqual(log.cabinet_name, 'Cabinet 1')
        self.assertEqual(station_one.status, CabinetSession.StatusChoices.OPEN)

        with override_settings(TAPTRACK_NFC_DEVICE_MAP=[{
            'device_id': 'CABINET1-STATION1',
            'api_key': self.device_api_key,
            'station': 'Station 1',
            'cabinet_name': 'Cabinet 1',
            'active': False,
        }]):
            disabled = self.client.post(
                '/api/verify-nfc/',
                {'nfc_uid': student.nfc_uid},
                HTTP_X_API_KEY=self.device_api_key,
                format='json',
            )
        self.assertEqual(disabled.status_code, 403)

    def test_development_fixture_supports_repeatable_nfc_enrollment(self):
        from django.core.management import call_command

        call_command('loaddata', 'taptrack_test_data', verbosity=0)
        call_command('loaddata', 'taptrack_test_data', verbosity=0)

        student = User.objects.get(username='taptrack_test_student')
        self.assertEqual(student.student_id, 'TAPTRACK-TEST-001')
        self.assertEqual(student.role, User.RoleChoices.STUDENT)
        self.assertIsNone(student.nfc_uid)
        self.assertFalse(student.has_usable_password())
        self.assertEqual(student.section.subject_code, 'TAP-TEST-2026')
        self.assertEqual(User.objects.filter(username='taptrack_test_student').count(), 1)
        self.assertEqual(Section.objects.filter(subject_code='TAP-TEST-2026').count(), 1)

        test_uid = '1268010402'
        initial_verification = self.client.post('/api/verify-nfc/', {
            'nfc_uid': test_uid,
        }, HTTP_X_API_KEY=self.device_api_key, format='json')
        self.assertEqual(initial_verification.status_code, 404, initial_verification.json())
        self.assertTrue(initial_verification.json()['registration_required'])
        token = parse_qs(urlparse(initial_verification.json()['registration_url']).query)['token'][0]

        registration = self.client.post('/api/cabinet/enrollment/register/', {
            'token': token,
            'student_id': student.student_id,
            'full_name': 'Jordan Santos',
            'email': student.email,
            'password': 'Fixture-Only-Password-123',
            'confirm_password': 'Fixture-Only-Password-123',
        }, format='json')
        self.assertEqual(registration.status_code, 201, registration.json())

        verification = self.client.post('/api/verify-nfc/', {
            'nfc_uid': test_uid,
        }, HTTP_X_API_KEY=self.device_api_key, format='json')
        self.assertEqual(verification.status_code, 200, verification.json())
        self.assertEqual(verification.json()['name'], 'Jordan Santos')
        self.assertEqual(verification.json()['student_id'], 'TAPTRACK-TEST-001')
        self.assertEqual(verification.json()['role'], 'student')

    def test_cabinet_registration_persists_student_and_card_in_django(self):
        response = self.client.post('/api/cabinet/register/', {
            'full_name': 'Cabinet Student',
            'student_id': 'CAB-001',
            'email': 'cabinet.student@example.com',
            'section': 'CAB-101',
            'nfc_uid': 'card-cabinet-001',
        }, HTTP_X_API_KEY=self.device_api_key, format='json')

        self.assertEqual(response.status_code, 201, response.json())
        user = User.objects.get(student_id='CAB-001')
        self.assertEqual(user.nfc_uid, 'CARD-CABINET-001')
        self.assertEqual(user.section, self.section)
        self.assertFalse(user.has_usable_password())

    def test_cabinet_registration_requires_device_api_key(self):
        response = self.client.post('/api/cabinet/register/', {
            'full_name': 'Cabinet Student',
            'student_id': 'CAB-002',
            'email': 'cabinet.student2@example.com',
            'nfc_uid': 'CARD-CABINET-002',
        }, format='json')

        self.assertEqual(response.status_code, 403)

    @override_settings(TAPTRACK_CABINET_NAME='Cabinet 1')
    def test_valid_cabinet_nfc_scan_creates_successful_access_log(self):
        student = User.objects.create_user(
            username='cabinet-access-student',
            email='cabinet-access@example.com',
            password='secret1234',
            role=User.RoleChoices.STUDENT,
            student_id='CAB-ACCESS-001',
            nfc_uid='CARD-CABINET-ACCESS-001',
        )
        session = CabinetSession.objects.create(station='Station 1', status=CabinetSession.StatusChoices.OPEN)
        session.opened_by.add(student)

        response = self.client.post('/api/verify-nfc/', {
            'nfc_uid': ' card- cabinet - access-001 ',
        }, HTTP_X_API_KEY=self.device_api_key, format='json')

        self.assertEqual(response.status_code, 200, response.json())
        log = student.access_logs.get()
        self.assertEqual(log.user, student)
        self.assertEqual(log.cabinet_session, session)
        self.assertIsNotNone(log.access_time)
        self.assertEqual(log.status, AccessLog.AccessStatusChoices.SUCCESS)
        self.assertEqual(log.action, AccessLog.AccessActionChoices.SCAN)
        self.assertEqual(log.nfc_uid, 'CARD-CABINET-ACCESS-001')
        self.assertEqual(log.station, 'Station 1')
        self.assertEqual(log.cabinet_name, 'Cabinet 1')
        self.assertEqual(log.reason, 'NFC verified successfully')
        self.assertIsNotNone(log.updated_at)
        self.assertEqual(response.json()['event_id'], log.id)
        self.assertEqual(response.json()['nfc_uid'], log.nfc_uid)
        self.assertEqual(response.json()['action'], log.action)
        self.assertEqual(response.json()['cabinet_session_id'], session.pk)
        self.assertEqual(response.json()['reason'], log.reason)

    @override_settings(
        TAPTRACK_NFC_DEVICE_MAP=[{'device_id': 'CABINET1', 'api_key': 'cabinet-device-test-key', 'station': 'Station 1', 'cabinet_name': 'Cabinet 1', 'active': True}],
        TAPTRACK_CABINET_STATIONS=['Station 1', 'Station 2'],
    )
    def test_single_reader_logs_the_selected_station_when_both_are_occupied(self):
        station_response = self.client.get(
            '/api/cabinet/workflow/',
            HTTP_X_API_KEY=self.device_api_key,
        )
        self.assertEqual(station_response.status_code, 200, station_response.json())
        self.assertEqual(station_response.json()['mode'], 'hardware')
        self.assertEqual(
            [(item['name'], item['status']) for item in station_response.json()['stations']],
            [('Station 1', 'available'), ('Station 2', 'available')],
        )

        student = User.objects.create_user(
            username='multi-station-student',
            email='multi-station@example.com',
            password='secret1234',
            role=User.RoleChoices.STUDENT,
            student_id='MULTI-STATION-001',
            nfc_uid='MULTI-STATION-UID',
        )
        station_one = CabinetSession.objects.create(station='Station 1', status=CabinetSession.StatusChoices.OPEN)
        station_two = CabinetSession.objects.create(
            station='Station 2',
            status=CabinetSession.StatusChoices.OPEN,
            workflow_state=CabinetSession.WorkflowStateChoices.OPENING,
        )
        station_one.opened_by.add(student)

        response = self.client.post(
            '/api/verify-nfc/',
            {'nfc_uid': student.nfc_uid},
            HTTP_X_API_KEY=self.device_api_key,
            format='json',
        )
        self.assertEqual(response.status_code, 200, response.json())

        logs = list(student.access_logs.order_by('station'))
        self.assertEqual([log.station for log in logs], ['Station 2'])
        self.assertEqual([log.cabinet_name for log in logs], ['Cabinet 1'])
        self.assertEqual([log.cabinet_session_id for log in logs], [station_two.pk])
        self.assertEqual(station_one.status, CabinetSession.StatusChoices.OPEN)

    @override_settings(
        TAPTRACK_NFC_DEVICE_MAP=[{'device_id': 'CABINET1', 'api_key': 'reader-workflow-secret', 'cabinet_name': 'Cabinet 1', 'active': True}],
        TAPTRACK_CABINET_STATIONS=['Station 1', 'Station 2'],
    )
    def test_backend_workflow_derives_open_close_and_rejected_audit_events(self):
        participant = User.objects.create_user(
            username='workflow-participant',
            email='workflow-participant@example.com',
            password='secret1234',
            role=User.RoleChoices.STUDENT,
            student_id='WORKFLOW-001',
            nfc_uid='WORKFLOW-UID-001',
        )
        outsider = User.objects.create_user(
            username='workflow-outsider',
            email='workflow-outsider@example.com',
            password='secret1234',
            role=User.RoleChoices.STUDENT,
            student_id='WORKFLOW-002',
            nfc_uid='WORKFLOW-UID-002',
        )

        opened = self.client.post(
            '/api/cabinet/workflow/', {'command': 'open', 'station': 'Station 1'},
            HTTP_X_API_KEY='reader-workflow-secret', format='json',
        )
        self.assertEqual(opened.status_code, 200, opened.json())
        session_id = opened.json()['session']['id']
        self.assertEqual(opened.json()['station'], 'Station 1')

        participant_scan = self.client.post(
            '/api/verify-nfc/', {'nfc_uid': participant.nfc_uid},
            HTTP_X_API_KEY='reader-workflow-secret', format='json',
        )
        self.assertEqual(participant_scan.status_code, 200, participant_scan.json())
        open_log = AccessLog.objects.get(user=participant)
        self.assertEqual(open_log.action, AccessLog.AccessActionChoices.OPEN)
        self.assertEqual(open_log.cabinet_session_id, session_id)

        finish_open = self.client.post(
            '/api/cabinet/workflow/', {'command': 'finish_open', 'station': 'Station 1'},
            HTTP_X_API_KEY='reader-workflow-secret', format='json',
        )
        self.assertEqual(finish_open.status_code, 200, finish_open.json())
        station_two_open = self.client.post(
            '/api/cabinet/workflow/', {'command': 'open', 'station': 'Station 2'},
            HTTP_X_API_KEY='reader-workflow-secret', format='json',
        )
        self.assertEqual(station_two_open.status_code, 200, station_two_open.json())
        station_two_scan = self.client.post(
            '/api/verify-nfc/', {'nfc_uid': participant.nfc_uid},
            HTTP_X_API_KEY='reader-workflow-secret', format='json',
        )
        self.assertEqual(station_two_scan.status_code, 200, station_two_scan.json())
        self.assertEqual(AccessLog.objects.filter(user=participant).latest('access_time').station, 'Station 2')
        station_two_finish = self.client.post(
            '/api/cabinet/workflow/', {'command': 'finish_open', 'station': 'Station 2'},
            HTTP_X_API_KEY='reader-workflow-secret', format='json',
        )
        self.assertEqual(station_two_finish.status_code, 200, station_two_finish.json())
        both_occupied = self.client.get(
            '/api/cabinet/workflow/', HTTP_X_API_KEY='reader-workflow-secret',
        )
        self.assertEqual(
            [(item['name'], item['status']) for item in both_occupied.json()['stations']],
            [('Station 1', 'occupied'), ('Station 2', 'occupied')],
        )
        self.client.post(
            '/api/cabinet/workflow/', {'command': 'start_close', 'station': 'Station 1'},
            HTTP_X_API_KEY='reader-workflow-secret', format='json',
        )

        rejected_scan = self.client.post(
            '/api/verify-nfc/', {'nfc_uid': outsider.nfc_uid},
            HTTP_X_API_KEY='reader-workflow-secret', format='json',
        )
        self.assertEqual(rejected_scan.status_code, 403, rejected_scan.json())
        rejected_log = AccessLog.objects.get(user=outsider)
        self.assertEqual(rejected_log.status, AccessLog.AccessStatusChoices.REJECTED)
        self.assertEqual(rejected_log.action, AccessLog.AccessActionChoices.CLOSE)
        self.assertEqual(rejected_log.reason, 'Participant is not part of this cabinet session')

        closing_scan = self.client.post(
            '/api/verify-nfc/', {'nfc_uid': participant.nfc_uid},
            HTTP_X_API_KEY='reader-workflow-secret', format='json',
        )
        self.assertEqual(closing_scan.status_code, 200, closing_scan.json())
        close_log = AccessLog.objects.filter(user=participant).latest('access_time')
        self.assertEqual(close_log.action, AccessLog.AccessActionChoices.CLOSE)
        self.assertEqual(close_log.reason, 'Cabinet close attendance recorded')

        closed = self.client.post(
            '/api/cabinet/workflow/', {'command': 'finish_close', 'station': 'Station 1'},
            HTTP_X_API_KEY='reader-workflow-secret', format='json',
        )
        self.assertEqual(closed.status_code, 200, closed.json())
        self.assertEqual(closed.json()['session']['status'], CabinetSession.StatusChoices.CLOSED)
        station_one_session = CabinetSession.objects.get(pk=session_id)
        station_two_session = CabinetSession.objects.get(pk=station_two_open.json()['session']['id'])
        self.assertEqual(station_one_session.status, CabinetSession.StatusChoices.CLOSED)
        self.assertEqual(station_two_session.status, CabinetSession.StatusChoices.OPEN)

        start_station_two_close = self.client.post(
            '/api/cabinet/workflow/', {'command': 'start_close', 'station': 'Station 2'},
            HTTP_X_API_KEY='reader-workflow-secret', format='json',
        )
        self.assertEqual(start_station_two_close.status_code, 200, start_station_two_close.json())
        station_two_closing_scan = self.client.post(
            '/api/verify-nfc/', {'nfc_uid': participant.nfc_uid},
            HTTP_X_API_KEY='reader-workflow-secret', format='json',
        )
        self.assertEqual(station_two_closing_scan.status_code, 200, station_two_closing_scan.json())
        self.assertEqual(station_two_closing_scan.json()['station']['name'], 'Station 2')
        finish_station_two_close = self.client.post(
            '/api/cabinet/workflow/', {'command': 'finish_close', 'station': 'Station 2'},
            HTTP_X_API_KEY='reader-workflow-secret', format='json',
        )
        self.assertEqual(finish_station_two_close.status_code, 200, finish_station_two_close.json())

        all_available = self.client.get(
            '/api/cabinet/workflow/', HTTP_X_API_KEY='reader-workflow-secret',
        )
        self.assertEqual(
            [(item['name'], item['status']) for item in all_available.json()['stations']],
            [('Station 1', 'available'), ('Station 2', 'available')],
        )

    def test_duplicate_participant_scan_is_recorded_as_duplicate(self):
        student = User.objects.create_user(
            username='cabinet-duplicate-student',
            email='cabinet-duplicate@example.com',
            password='secret1234',
            role=User.RoleChoices.STUDENT,
            student_id='CAB-ACCESS-DUPLICATE',
            nfc_uid='CARD-CABINET-DUPLICATE',
        )
        session = CabinetSession.objects.create(station='Station 1', status=CabinetSession.StatusChoices.OPEN)
        session.opened_by.add(student)

        first = self.client.post('/api/verify-nfc/', {'nfc_uid': student.nfc_uid}, HTTP_X_API_KEY=self.device_api_key, format='json')
        duplicate = self.client.post('/api/verify-nfc/', {'nfc_uid': student.nfc_uid}, HTTP_X_API_KEY=self.device_api_key, format='json')

        self.assertEqual(first.status_code, 200, first.json())
        self.assertEqual(duplicate.status_code, 409, duplicate.json())
        log = student.access_logs.order_by('-access_time').first()
        self.assertEqual(log.status, AccessLog.AccessStatusChoices.DUPLICATE)
        self.assertEqual(log.reason, 'Duplicate participant scan')
        self.assertEqual(log.cabinet_session, session)
        self.assertEqual(log.station, 'Station 1')
        event = CabinetEvent.objects.get(details__access_log_id=log.pk)
        self.assertEqual(event.user_id, student.pk)
        self.assertEqual(event.event_type, 'access_denied')

    def test_multiple_valid_cabinet_scans_create_multiple_access_logs(self):
        student = User.objects.create_user(
            username='cabinet-repeat-student',
            email='cabinet-repeat@example.com',
            password='secret1234',
            role=User.RoleChoices.STUDENT,
            student_id='CAB-ACCESS-002',
            nfc_uid='CARD-CABINET-ACCESS-002',
        )

        for _ in range(2):
            response = self.client.post('/api/verify-nfc/', {
                'nfc_uid': student.nfc_uid,
            }, HTTP_X_API_KEY=self.device_api_key, format='json')
            self.assertEqual(response.status_code, 200, response.json())

        self.assertEqual(student.access_logs.filter(status='success').count(), 2)
        events = list(student.cabinet_events.order_by('timestamp'))
        self.assertEqual(len(events), 2)
        self.assertEqual({event.user_id for event in events}, {student.pk})
        self.assertEqual({event.event_type for event in events}, {'access_granted'})
        for event in events:
            self.assertTrue(student.access_logs.filter(pk=event.details['access_log_id']).exists())

    def test_invalid_cabinet_nfc_scan_creates_no_success_log(self):
        response = self.client.post('/api/verify-nfc/', {
            'nfc_uid': 'CARD-UNKNOWN-999',
        }, HTTP_X_API_KEY=self.device_api_key, format='json')

        self.assertEqual(response.status_code, 404, response.json())
        self.assertFalse(AccessLog.objects.filter(status='success').exists())
        log = AccessLog.objects.get()
        self.assertIsNone(log.user)
        self.assertEqual(log.nfc_uid, 'CARD-UNKNOWN-999')
        self.assertEqual(log.status, AccessLog.AccessStatusChoices.UNREGISTERED)
        self.assertEqual(log.action, AccessLog.AccessActionChoices.SCAN)
        self.assertEqual(log.reason, 'NFC UID is not registered')
        self.assertFalse(CabinetEvent.objects.exists())

    def test_unknown_nfc_scan_creates_temporary_enrollment_session(self):
        response = self.client.post('/api/verify-nfc/', {
            'nfc_uid': 'card-enrollment-001',
        }, HTTP_X_API_KEY=self.device_api_key, format='json')

        self.assertEqual(response.status_code, 404, response.json())
        payload = response.json()
        self.assertFalse(payload['registered'])
        self.assertTrue(payload['registration_required'])
        self.assertNotIn('card-enrollment-001', payload['registration_url'])

        enrollment = NFCEnrollmentSession.objects.get(nfc_uid='CARD-ENROLLMENT-001')
        token = parse_qs(urlparse(payload['registration_url']).query)['token'][0]
        validation = self.client.get('/api/cabinet/enrollment/validate/', {'token': token})

        self.assertEqual(validation.status_code, 200, validation.json())
        self.assertTrue(validation.json()['valid'])
        self.assertTrue(enrollment.verify_token(token))

    def test_repeated_unknown_nfc_scans_reuse_one_pending_session(self):
        first = self.client.post('/api/verify-nfc/', {
            'nfc_uid': 'CARD-REUSE-001',
        }, HTTP_X_API_KEY=self.device_api_key, format='json')
        second = self.client.post('/api/verify-nfc/', {
            'nfc_uid': 'CARD-REUSE-001',
        }, HTTP_X_API_KEY=self.device_api_key, format='json')

        self.assertEqual(first.status_code, 404, first.json())
        self.assertEqual(second.status_code, 404, second.json())
        self.assertEqual(
            NFCEnrollmentSession.objects.filter(
                nfc_uid='CARD-REUSE-001',
                status=NFCEnrollmentSession.StatusChoices.PENDING,
            ).count(),
            1,
        )
        first_token = parse_qs(urlparse(first.json()['registration_url']).query)['token'][0]
        second_token = parse_qs(urlparse(second.json()['registration_url']).query)['token'][0]
        self.assertNotEqual(first_token, second_token)
        self.assertEqual(self.client.get('/api/cabinet/enrollment/validate/', {'token': first_token}).status_code, 410)
        self.assertEqual(self.client.get('/api/cabinet/enrollment/validate/', {'token': second_token}).status_code, 200)

    def test_invalid_expired_and_completed_enrollment_tokens_are_rejected(self):
        invalid = self.client.get('/api/cabinet/enrollment/validate/', {'token': 'invalid-token'})
        self.assertEqual(invalid.status_code, 410)

        enrollment = NFCEnrollmentSession(
            nfc_uid='CARD-ENROLLMENT-002',
            token_digest='',
            token_hash='',
            expires_at=timezone.now() - timedelta(minutes=1),
        )
        enrollment.set_token('expired-token')
        enrollment.save()
        expired = self.client.get('/api/cabinet/enrollment/validate/', {'token': 'expired-token'})
        self.assertEqual(expired.status_code, 410)
        enrollment.refresh_from_db()
        self.assertEqual(enrollment.status, NFCEnrollmentSession.StatusChoices.EXPIRED)

        enrollment = NFCEnrollmentSession(
            nfc_uid='CARD-ENROLLMENT-003',
            token_digest='',
            token_hash='',
            expires_at=timezone.now() + timedelta(minutes=15),
            status=NFCEnrollmentSession.StatusChoices.COMPLETED,
        )
        enrollment.set_token('completed-token')
        enrollment.save()
        completed = self.client.get('/api/cabinet/enrollment/validate/', {'token': 'completed-token'})
        self.assertEqual(completed.status_code, 410)

    def test_pending_student_registration_links_nfc_and_consumes_token(self):
        student = User.objects.create_user(
            username='pending-registration-student',
            email='pending-registration@example.com',
            password=None,
            role=User.RoleChoices.STUDENT,
            student_id='PENDING-001',
            nfc_uid=None,
            is_active=True,
        )
        enrollment = NFCEnrollmentSession(
            nfc_uid='CARD-ENROLLMENT-READY',
            token_digest='',
            token_hash='',
            expires_at=timezone.now() + timedelta(minutes=15),
        )
        enrollment.set_token('registration-token')
        enrollment.save()

        response = self.client.post('/api/cabinet/enrollment/register/', {
            'token': 'registration-token',
            'student_id': 'PENDING-001',
            'full_name': 'Pending Student',
            'email': 'pending.registration@example.com',
            'password': 'new-password-123',
            'confirm_password': 'new-password-123',
        }, format='json')

        self.assertEqual(response.status_code, 201, response.json())
        student.refresh_from_db()
        enrollment.refresh_from_db()
        self.assertEqual(student.nfc_uid, 'CARD-ENROLLMENT-READY')
        self.assertTrue(student.has_usable_password())
        self.assertEqual(enrollment.status, NFCEnrollmentSession.StatusChoices.COMPLETED)

        reused = self.client.post('/api/cabinet/enrollment/register/', {
            'token': 'registration-token',
            'student_id': 'PENDING-001',
            'full_name': 'Pending Student',
            'email': 'pending.registration@example.com',
            'password': 'new-password-123',
            'confirm_password': 'new-password-123',
        }, format='json')
        self.assertEqual(reused.status_code, 410, reused.json())

    def test_enrollment_registration_rejects_unknown_student_id(self):
        enrollment = NFCEnrollmentSession(
            nfc_uid='CARD-ENROLLMENT-UNKNOWN-STUDENT',
            token_digest='',
            token_hash='',
            expires_at=timezone.now() + timedelta(minutes=15),
        )
        enrollment.set_token('unknown-student-token')
        enrollment.save()

        response = self.client.post('/api/cabinet/enrollment/register/', {
            'token': 'unknown-student-token',
            'student_id': 'DOES-NOT-EXIST',
            'full_name': 'Unknown Student',
            'email': 'unknown.student@example.com',
            'password': 'new-password-123',
            'confirm_password': 'new-password-123',
        }, format='json')

        self.assertEqual(response.status_code, 404, response.json())

    def test_enrollment_registration_rejects_existing_nfc_owner(self):
        User.objects.create_user(
            username='other-nfc-owner',
            email='other-nfc-owner@example.com',
            password='existing-password',
            role=User.RoleChoices.STUDENT,
            student_id='OWNER-001',
            nfc_uid='CARD-ALREADY-ASSIGNED',
        )
        User.objects.create_user(
            username='pending-nfc-student',
            email='pending-nfc-student@example.com',
            password=None,
            role=User.RoleChoices.STUDENT,
            student_id='PENDING-002',
            nfc_uid=None,
        )
        enrollment = NFCEnrollmentSession(
            nfc_uid='CARD-ALREADY-ASSIGNED',
            token_digest='',
            token_hash='',
            expires_at=timezone.now() + timedelta(minutes=15),
        )
        enrollment.set_token('assigned-token')
        enrollment.save()

        response = self.client.post('/api/cabinet/enrollment/register/', {
            'token': 'assigned-token',
            'student_id': 'PENDING-002',
            'full_name': 'Pending NFC Student',
            'email': 'pending.nfc@example.com',
            'password': 'new-password-123',
            'confirm_password': 'new-password-123',
        }, format='json')
        self.assertEqual(response.status_code, 409, response.json())

    def test_enrollment_registration_rejects_existing_account(self):
        User.objects.create_user(
            username='existing-account-student',
            email='existing-account@example.com',
            password='existing-password',
            role=User.RoleChoices.STUDENT,
            student_id='EXISTING-001',
            nfc_uid=None,
        )
        enrollment = NFCEnrollmentSession(
            nfc_uid='CARD-EXISTING-ACCOUNT',
            token_digest='',
            token_hash='',
            expires_at=timezone.now() + timedelta(minutes=15),
        )
        enrollment.set_token('existing-account-token')
        enrollment.save()

        response = self.client.post('/api/cabinet/enrollment/register/', {
            'token': 'existing-account-token',
            'student_id': 'EXISTING-001',
            'full_name': 'Existing Account Student',
            'email': 'existing.account@example.com',
            'password': 'new-password-123',
            'confirm_password': 'new-password-123',
        }, format='json')
        self.assertEqual(response.status_code, 409, response.json())
